"""Excel import of the existing materials list: products, suppliers and product photos.

Workflow (spec section 23): select the workbook -> select the worksheet -> map the columns ->
preview -> validate -> resolve duplicates -> import. Nothing reaches the database before the
person has seen the preview, and the import itself:

* takes a safety backup first (section 30) and stops if that fails;
* writes every product, supplier, category and photo in ONE transaction - all or nothing;
* copies every photo into the WMS attachment store (section 24), so the workbook and any picture
  folder can be deleted afterwards.

Photos are found in four ways (section 24): pictures embedded in the sheet (floating over a cell or
placed in a cell with Excel's "Place in Cell"), a full file path in a cell, a hyperlink to a file,
or just a file name looked up in a chosen picture folder. Web addresses are reported, not
downloaded: the warehouse PC may be offline and a web page is not a picture.
"""

from __future__ import annotations

import json
import posixpath
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import delete, func, select

from wms.core.enums import CLOSED_SITE_STATUSES, EntityType, ItemType, SiteStatus
from wms.core.exceptions import ValidationError, WMSError
from wms.core.numbers import ZERO
from wms.core.permissions import P
from wms.db.models import (
    Attachment,
    Bin,
    Category,
    Product,
    ProjectMaterial,
    SavedSearch,
    Site,
    Supplier,
    Unit,
)
from wms.services import audit
from wms.services.attachments import PHOTO_EXT, AttachmentService, StagedFile
from wms.services.base import BaseService, normalise_code
from wms.services.products import ProductService, code_key

# --- what can be imported ------------------------------------------------------------------

#: field -> (label, header words that suggest it). Order = order on the mapping screen.
FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "internal_code": ("Code *", ("code", "internal code", "product code", "item code", "our code", "sku",
                                 "material code", "stock code", "код")),
    "name": ("Name *", ("name", "product name", "item name", "material", "material name", "product", "назва")),
    "description": ("Description", ("description", "desc", "details", "опис", "name / description")),
    "supplier_code": ("Product Number (supplier's code)", (
        "product number", "product no", "product no.", "supplier code", "supplier product code", "part number",
        "part no", "part no.", "article", "article number", "ref", "reference", "decor", "decor code", "артикул")),
    "supplier": ("Supplier", ("supplier", "vendor", "supplier name", "постачальник")),
    "barcode": ("Barcode", ("barcode", "ean", "ean13", "gtin", "upc", "штрихкод")),
    "brand": ("Brand", ("brand", "manufacturer", "make", "бренд")),
    "category": ("Category", ("category", "group", "product group", "категорія")),
    "unit": ("Unit", ("unit", "uom", "unit of measure", "units", "од")),
    "min_stock": ("Minimum stock", ("min", "minimum", "min stock", "minimum stock", "min level")),
    "reorder_level": ("Reorder level", ("reorder level", "reorder point")),
    "reorder_qty": ("Reorder quantity", ("reorder qty", "reorder quantity", "order qty")),
    "unit_cost": ("Unit cost (£)", ("cost", "unit cost", "price", "unit price", "cost price", "£", "ціна")),
    "default_bin": ("Default bin", ("bin", "location", "default bin", "комірка")),
    "photo": ("Photo", ("photo", "image", "picture", "pic", "product photo", "фото", "img")),
    "notes": ("Notes", ("notes", "note", "comment", "comments", "примітки")),
    # project material list (spec section 34): never stock, only "what the project needs"
    "project": ("Project", ("project", "project number", "project no", "project no.", "job", "job number",
                            "job no", "site", "проєкт", "проект")),
    "required_qty": ("Required Quantity", ("required", "required qty", "required quantity", "qty required",
                                           "quantity required", "needed", "need", "потрібно")),
}
REQUIRED = ("internal_code", "name")
QTY_FIELDS = ("min_stock", "reorder_level", "reorder_qty")

CREATE, UPDATE, SKIP = "CREATE", "UPDATE", "SKIP"
#: LINK = "Match": use the suggested existing material as it is (UI refactor spec 31) - on a project
#: import its Required quantity goes on that material; nothing about the material changes
LINK = "LINK"
ACTION_LABELS = {CREATE: "New", UPDATE: "Update existing", SKIP: "Skip", LINK: "Match (use existing)"}
#: how an existing material was found: CODE is certain; the others are suggestions the person decides
UNCERTAIN = ("PRODUCT_NO", "SIMILAR_CODE")

#: what the preview shows per row (spec section 30)
NEW, DUPLICATE, ERROR, SKIPPED = "NEW", "DUPLICATE", "ERROR", "SKIPPED"
MATCHED = "MATCHED"  # a suggested existing material the person chose to use
STATUS_LABELS = {NEW: "NEW", UPDATE: "UPDATE", DUPLICATE: "DUPLICATE", ERROR: "ERROR", SKIPPED: "SKIPPED"}

#: how an unknown supplier / project in the file is resolved (spec sections 35, 37)
MATCH, IGNORE, SKIP_ROW = "MATCH", "IGNORE", "SKIP_ROW"
SUPPLIER_ACTIONS = {MATCH: "Match to an existing supplier", CREATE: "Create a new supplier",
                    IGNORE: "Ignore (leave the supplier empty)"}
PROJECT_ACTIONS = {MATCH: "Match to an existing project", CREATE: "Create a new project",
                   IGNORE: "Ignore (import the material, no project)", SKIP_ROW: "Skip these rows"}

PROFILE_PREFIX = '{"import_profile":'  # mapping profiles live in saved_searches, like saved views

_COMPANY_WORDS = {"LTD", "LIMITED", "PLC", "UK", "CO", "COMPANY", "GMBH", "INC", "LLC", "SA", "AG", "BV", "NV",
                  "SRL", "SPA", "GROUP", "THE", "AND", "INTERNATIONAL", "INT", "EUROPE"}


def supplier_key(name: str) -> str:
    """'Egger UK Ltd', 'EGGER', 'egger (uk) limited' -> 'EGGER': case, punctuation and company
    words do not make a different supplier."""
    words = re.sub(r"[^A-Z0-9]+", " ", (name or "").upper()).split()
    core = [w for w in words if w not in _COMPANY_WORDS]
    return "".join(core or words)

MAX_HEADER_SCAN = 15


# --- results of looking at a workbook ----------------------------------------------------------

@dataclass
class SheetInfo:
    name: str
    rows: int
    header_row: int  # 1-based
    headers: list[str]  # column titles of the header row ("Column C" when blank)
    samples: list[list[str]]  # first rows under the header, as text
    guess: dict[str, int]  # field -> 0-based column
    floating_images: int = 0
    cell_images: int = 0
    hyperlinks: int = 0
    paths: int = 0
    filenames: int = 0
    web_links: int = 0
    profile: str | None = None  # the saved mapping profile used for ``guess``

    @property
    def photo_summary(self) -> str:
        parts = [f"{n} {what}" for n, what in (
            (self.floating_images, "picture(s) embedded over cells"), (self.cell_images, "picture(s) placed in cells"),
            (self.hyperlinks, "hyperlink(s) to files"), (self.paths, "file path(s)"),
            (self.filenames, "file name(s)"), (self.web_links, "web address(es) - not downloaded")) if n]
        return ", ".join(parts) if parts else "No photos found"


@dataclass
class WorkbookInfo:
    path: Path
    sheets: list[SheetInfo]

    def sheet(self, name: str) -> SheetInfo:
        return next(s for s in self.sheets if s.name == name)


@dataclass
class ImportOptions:
    sheet: str
    header_row: int
    mapping: dict[str, int]  # field -> 0-based column
    photo_folder: Path | None = None
    default_unit_id: int | None = None
    item_type: str = ItemType.RAW_MATERIAL.value
    on_existing: str = SKIP  # what to do with codes already in the WMS: SKIP / UPDATE
    create_suppliers: bool = True  # default for unknown suppliers: Create (else Ignore) - shown to confirm
    create_categories: bool = True
    project_id: int | None = None  # "Import into project": every row goes on this project's list


@dataclass
class PhotoSource:
    kind: str  # EMBEDDED / IN_CELL / PATH / HYPERLINK / FILENAME
    label: str  # what the person sees ("embedded picture", the path, the file name)
    path: Path | None = None  # file on disk (paths, hyperlinks, file names)
    data: bytes | None = None  # embedded picture bytes
    ext: str = ".png"


@dataclass
class ImportRow:
    row: int  # Excel row number
    values: dict[str, str]
    action: str = CREATE
    existing_id: int | None = None
    existing_label: str = ""
    match_kind: str | None = None  # CODE / PRODUCT_NO - how the existing material was found
    photo: PhotoSource | None = None
    photo_missing: bool = False  # the file names a photo that could not be used
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    changes: list[tuple[str, str, str]] = field(default_factory=list)  # (field, in the WMS, in Excel)
    duplicate_of: int | None = None  # the same code is on an earlier row of the file
    project_key: str | None = None  # key into ImportPlan.projects
    required: Decimal | None = None  # required quantity for the project list

    @property
    def ok(self) -> bool:
        return not self.errors

    @property
    def status(self) -> str:
        if self.errors:
            return ERROR
        if self.duplicate_of is not None:
            return DUPLICATE
        if self.action == CREATE:
            return NEW
        if self.action == UPDATE:
            return UPDATE
        if self.action == LINK:
            return MATCHED
        if self.existing_id is not None and self.match_kind is not None and not self.archived:
            return DUPLICATE
        return SKIPPED

    @property
    def archived(self) -> bool:
        return any("removed from active use" in w for w in self.warnings)

    @property
    def changes_text(self) -> str:
        if self.action != UPDATE and self.status != DUPLICATE:
            return ""
        if not self.changes:
            return "no changes" if self.existing_id else ""
        return "; ".join(f"{f}: {old or '(empty)'} → {new}" for f, old, new in self.changes)

    @property
    def code(self) -> str:
        return self.values.get("internal_code", "")

    @property
    def problems(self) -> str:
        return "; ".join(self.errors + self.warnings)


@dataclass
class SupplierChoice:
    """A supplier name from the file that is not exactly a WMS supplier (spec section 37)."""
    text: str  # as written in the file (first occurrence)
    rows: int = 0
    action: str = CREATE
    match_id: int | None = None
    match_name: str = ""
    suggested: bool = False  # match_id came from the normalised name (EGGER = Egger UK Ltd)


@dataclass
class ProjectChoice:
    """A project named in the file's Project column (spec section 35)."""
    text: str
    rows: int = 0
    action: str = IGNORE
    match_id: int | None = None
    match_label: str = ""
    closed: bool = False  # the matched project is completed / cancelled / archived


@dataclass
class ImportPlan:
    source: Path
    options: ImportOptions
    rows: list[ImportRow]
    suppliers: dict[str, SupplierChoice] = field(default_factory=dict)  # supplier_key -> choice
    projects: dict[str, ProjectChoice] = field(default_factory=dict)  # upper text -> choice
    project_label: str = ""  # "Import into project" target
    #: Compare with Excel (spec 33): rows the person unticked, and list lines to take off the project
    excluded: set[int] = field(default_factory=set)  # Excel row numbers
    remove_ids: list[int] = field(default_factory=list)  # product ids (project import only)

    @property
    def has_project(self) -> bool:
        return bool(self.options.project_id) or "project" in self.options.mapping

    def skipped_by_project(self, r: ImportRow) -> bool:
        c = self.projects.get(r.project_key) if r.project_key else None
        return c is not None and c.action == SKIP_ROW

    def links_project(self, r: ImportRow, respect_excluded: bool = True) -> bool:
        """The row puts its material on a project's list."""
        if not r.ok or r.duplicate_of is not None or r.required is None or self.skipped_by_project(r) \
                or (respect_excluded and r.row in self.excluded):
            return False
        if r.match_kind in UNCERTAIN:
            # a suggested match is used only when the person said so (Match / Update / New)
            material = r.action in (LINK, UPDATE, CREATE)
        else:
            material = r.action != SKIP or r.existing_id is not None
        if self.options.project_id:
            return material
        c = self.projects.get(r.project_key) if r.project_key else None
        return c is not None and c.action in (MATCH, CREATE) and material

    @property
    def to_import(self) -> list[ImportRow]:
        """Rows that change something: a new / updated material, or a project list line."""
        return [r for r in self.rows if r.ok and r.duplicate_of is None and not self.skipped_by_project(r)
                and r.row not in self.excluded and (r.action in (CREATE, UPDATE) or self.links_project(r))]

    def counts(self) -> dict[str, int]:
        rows = [r for r in self.rows if not self.skipped_by_project(r)]
        live = [r for r in rows if r.ok and r.action in (CREATE, UPDATE) and r.row not in self.excluded]
        return {"rows": len(self.rows), "create": sum(1 for r in live if r.action == CREATE),
                "update": sum(1 for r in live if r.action == UPDATE),
                "duplicate": sum(1 for r in rows if r.status == DUPLICATE),
                "skip": sum(1 for r in rows if r.status == SKIPPED) + (len(self.rows) - len(rows)),
                "errors": sum(1 for r in rows if not r.ok),
                "photos": sum(1 for r in live if r.photo is not None),
                "photos_missing": sum(1 for r in rows if r.photo_missing),
                "project_lines": sum(1 for r in rows if self.links_project(r)),
                "matched": sum(1 for r in rows if r.ok and r.action == LINK),
                "remove": len(self.remove_ids)}

    def summary(self) -> str:
        c = self.counts()
        text = (f"{c['rows']} row(s): {c['create']} new, {c['update']} to update, {c['duplicate']} already in "
                f"the WMS / duplicate, {c['skip']} skipped, {c['errors']} with errors (not imported), "
                f"{c['photos']} photo(s)")
        if c["project_lines"] or self.options.project_id or self.projects:
            text += f", {c['project_lines']} project material line(s)"
        return text

    def apply_to_similar(self, r: ImportRow, action: str) -> int:
        """"Apply to similar" (spec section 31): the same choice for every row in the same situation."""
        n = 0
        for other in self.rows:
            # similar = found the same way (new / same code / same product number) and free to choose
            if other.ok and other.duplicate_of is None and other.match_kind == r.match_kind \
                    and self.choices(other) == self.choices(r) and action in self.choices(other):
                other.action = action
                n += 1
        return n

    def choices(self, r: ImportRow) -> list[str]:
        if not r.ok or r.duplicate_of is not None or r.archived:
            return [SKIP]
        if r.match_kind == "CODE":
            return [UPDATE, SKIP]
        if r.match_kind in UNCERTAIN:
            # spec 31: [Match] [Create New] [Skip] - and Update when the file should correct it
            return [SKIP, LINK, UPDATE, CREATE] if self.has_project else [SKIP, UPDATE, CREATE]
        return [CREATE, SKIP]


ADDED, REMOVED, QTY_CHANGED, DETAIL_CHANGED, UNCHANGED = "ADDED", "REMOVED", "QTY", "DETAIL", "UNCHANGED"
COMPARE_LABELS = {ADDED: "Added", REMOVED: "Removed", QTY_CHANGED: "Quantity changed",
                  DETAIL_CHANGED: "Detail changed", UNCHANGED: "Unchanged"}


@dataclass
class CompareLine:
    """One line of "Compare with Excel": the project's current list against the revised file (spec 33)."""

    change: str  # ADDED / REMOVED / QTY / DETAIL / UNCHANGED
    code: str
    name: str
    current: Decimal | None  # Required on the project now (None = not on the list)
    new: Decimal | None  # Required in the file (None = not in the file)
    excel_row: int | None = None
    product_id: int | None = None
    details: str = ""

    @property
    def label(self) -> str:
        return COMPARE_LABELS[self.change]

    @property
    def difference(self) -> Decimal | None:
        if self.current is None and self.new is None:
            return None
        return (self.new or ZERO) - (self.current or ZERO)


@dataclass
class ImportResult:
    created: int
    updated: int
    skipped: int
    failed_rows: int
    photos: int
    suppliers_created: list[str]
    categories_created: list[str]
    backup: Path | None
    photo_problems: list[str] = field(default_factory=list)
    duplicates: int = 0
    photos_missing: int = 0
    project_lines: int = 0
    projects_created: list[str] = field(default_factory=list)
    project_lines_removed: int = 0
    total_rows: int = 0


# --- reading cells ----------------------------------------------------------------------------

def cell_text(value) -> str:
    """A cell as the text a person would type: 3311.0 -> '3311', dates as dd/mm/yyyy."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    return " ".join(str(value).split())


def _norm_header(text: str) -> str:
    return re.sub(r"[^a-z0-9£а-яіїєґ /]+", " ", text.lower()).strip()


def guess_mapping(headers: list[str]) -> dict[str, int]:
    """Field -> column for headers that obviously name a field (exact words first, then 'contains')."""
    out: dict[str, int] = {}
    used: set[int] = set()
    normed = [_norm_header(h) for h in headers]
    for exact in (True, False):
        for f, (_label, words) in FIELDS.items():
            if f in out:
                continue
            for i, h in enumerate(normed):
                if i in used or not h:
                    continue
                if (h in words) if exact else any(len(w) > 3 and w in h for w in words):
                    out[f] = i
                    used.add(i)
                    break
    # "Name / Description" in one column: it is the name
    if "name" not in out and "description" in out:
        out["name"] = out.pop("description")
    return out


def _is_header(values: list) -> bool:
    texts = [cell_text(v) for v in values if v not in (None, "")]
    return len(texts) >= 2 and len(guess_mapping(texts)) >= 1 and all(not isinstance(v, (int, float))
                                                                      for v in values if v not in (None, ""))


def _classify(text: str) -> str | None:
    t = text.strip()
    if not t:
        return None
    low = t.lower()
    if low.startswith(("http://", "https://", "www.")):
        return "WEB"
    if low.startswith("file:"):
        return "PATH"
    ext = Path(t.replace("\\", "/")).suffix.lower()
    if ext not in PHOTO_EXT:
        return None
    if "/" in t or "\\" in t:
        return "PATH"
    return "FILENAME"


# --- pictures inside the workbook ---------------------------------------------------------------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, str]:
    """rId -> target path (inside the zip) for a part's relationships."""
    folder, name = posixpath.split(part)
    rel_part = posixpath.join(folder, "_rels", name + ".rels")
    if rel_part not in z.namelist():
        return {}
    out = {}
    for el in ET.fromstring(z.read(rel_part)):
        target = el.get("Target", "")
        out[el.get("Id")] = (target.lstrip("/") if target.startswith("/")
                             else posixpath.normpath(posixpath.join(folder, target)))
    return out


def _sheet_parts(z: zipfile.ZipFile) -> dict[str, str]:
    """Worksheet name -> its part inside the zip."""
    rels = _rels(z, "xl/workbook.xml")
    out = {}
    for el in ET.fromstring(z.read("xl/workbook.xml")).iter():
        if _local(el.tag) == "sheet":
            rid = next((v for k, v in el.attrib.items() if _local(k) == "id"), None)
            if rid in rels:
                out[el.get("name")] = rels[rid]
    return out


def in_cell_pictures(path: Path) -> dict[str, dict[tuple[int, int], tuple[bytes, str]]]:
    """Pictures placed IN cells (Excel 365 "Place in Cell" / IMAGE()): sheet -> (row, col) ->
    (bytes, extension). openpyxl does not read these, so the rich-value parts are followed by hand:
    cell vm -> valueMetadata -> futureMetadata rvb -> rich value -> richValueRel -> media file."""
    out: dict[str, dict[tuple[int, int], tuple[bytes, str]]] = {}
    try:
        z = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        return out
    with z:
        names = set(z.namelist())
        rv_part, rel_part, meta_part = ("xl/richData/rdrichvalue.xml", "xl/richData/richValueRel.xml",
                                        "xl/metadata.xml")
        if not {rv_part, rel_part, meta_part} <= names:
            return out
        meta = ET.fromstring(z.read(meta_part))
        future: list[int | None] = []  # futureMetadata XLRICHVALUE bk -> rich value index
        value_meta: list[int | None] = []  # valueMetadata bk (vm - 1) -> futureMetadata index
        types = [el.get("name") for el in meta.iter() if _local(el.tag) == "metadataType"]
        for el in meta:
            if _local(el.tag) == "futureMetadata" and el.get("name") == "XLRICHVALUE":
                for bk in el:
                    rvb = next((x for x in bk.iter() if _local(x.tag) == "rvb"), None)
                    future.append(int(rvb.get("i")) if rvb is not None else None)
            elif _local(el.tag) == "valueMetadata":
                for bk in el:
                    rc = next((x for x in bk if _local(x.tag) == "rc"), None)
                    ok = rc is not None and 0 < int(rc.get("t", 0)) <= len(types) and \
                        types[int(rc.get("t")) - 1] == "XLRICHVALUE"
                    value_meta.append(int(rc.get("v")) if ok else None)
        # rich values: which "v" is the picture (structure key _rvRel:LocalImageIdentifier)
        key_pos: dict[int, int] = {}
        struct_part = "xl/richData/rdrichvaluestructure.xml"
        if struct_part in names:
            for i, st in enumerate(x for x in ET.fromstring(z.read(struct_part)) if _local(x.tag) == "s"):
                keys = [k.get("n") for k in st if _local(k.tag) == "k"]
                if "_rvRel:LocalImageIdentifier" in keys:
                    key_pos[i] = keys.index("_rvRel:LocalImageIdentifier")
        rich: list[int | None] = []
        for rv in (x for x in ET.fromstring(z.read(rv_part)) if _local(x.tag) == "rv"):
            vals = [v.text for v in rv if _local(v.tag) == "v"]
            pos = key_pos.get(int(rv.get("s", 0)), 0)
            try:
                rich.append(int(vals[pos]))
            except (IndexError, TypeError, ValueError):
                rich.append(None)
        rel_ids = [next((v for k, v in el.attrib.items() if _local(k) == "id"), None)
                   for el in ET.fromstring(z.read(rel_part)) if _local(el.tag) == "rel"]
        targets = _rels(z, rel_part)

        def picture(vm: int) -> tuple[bytes, str] | None:
            try:
                f = value_meta[vm - 1]
                r = future[f] if f is not None else None
                rel = rich[r] if r is not None else None
                target = targets.get(rel_ids[rel]) if rel is not None else None
            except IndexError:
                return None
            if not target or target not in names:
                return None
            return z.read(target), posixpath.splitext(target)[1].lower() or ".png"

        for sheet, part in _sheet_parts(z).items():
            if part not in names:
                continue
            found: dict[tuple[int, int], tuple[bytes, str]] = {}
            for _ev, el in ET.iterparse(z.open(part)):
                if _local(el.tag) == "c" and el.get("vm"):
                    pic = picture(int(el.get("vm")))
                    if pic:
                        row, col = _ref(el.get("r", ""))
                        found[(row, col)] = pic
                    el.clear()
            if found:
                out[sheet] = found
    return out


def _ref(ref: str) -> tuple[int, int]:
    """'E12' -> (12, 4): Excel row, 0-based column."""
    m = re.match(r"([A-Z]+)(\d+)", ref)
    if not m:
        return 0, 0
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col - 1


def floating_pictures(ws) -> dict[int, list[tuple[int, bytes, str]]]:
    """Pictures floating over the sheet, by the row of their top-left corner: row -> [(col, bytes, ext)]."""
    out: dict[int, list[tuple[int, bytes, str]]] = {}
    for img in getattr(ws, "_images", []):
        start = getattr(getattr(img, "anchor", None), "_from", None)
        if start is None:
            continue  # absolute position - no row to belong to
        try:
            data = img._data()
        except Exception:  # noqa: BLE001 - one unreadable picture must not stop the import
            continue
        ext = "." + (getattr(img, "format", None) or "png").lower().replace("jpeg", "jpg")
        out.setdefault(start.row + 1, []).append((start.col, data, ext))
    return out


# --- the service ------------------------------------------------------------------------------

class ExcelImportService(BaseService):

    def inspect(self, path: str | Path) -> WorkbookInfo:
        """Look at every worksheet: header row, columns, a few sample rows, the guessed mapping and
        how the photos are stored. Nothing is changed."""
        self.require(P.PRODUCTS_EDIT)
        path = self._check_file(path)
        wb = self._load(path)
        cells = in_cell_pictures(path)
        sheets = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN, values_only=True))
            header_row = next((i + 1 for i, r in enumerate(rows) if _is_header(list(r))), None)
            if header_row is None:
                header_row = next((i + 1 for i, r in enumerate(rows) if any(v not in (None, "") for v in r)), 1)
            headers = self._headers(ws, header_row)
            guess = guess_mapping(headers)
            samples = self._samples(ws, header_row, len(headers))
            info = SheetInfo(ws.title, max(ws.max_row - header_row, 0), header_row, headers, samples, guess,
                             floating_images=sum(len(v) for v in floating_pictures(ws).values()),
                             cell_images=len(cells.get(ws.title, {})))
            for row in ws.iter_rows(min_row=header_row + 1):
                for c in row:
                    if c.hyperlink is not None and c.hyperlink.target:
                        kind = _classify(c.hyperlink.target)
                        if kind == "WEB":
                            info.web_links += 1
                        elif kind:
                            info.hyperlinks += 1
                        continue
                    if isinstance(c.value, str):
                        kind = _classify(c.value)
                        if kind == "PATH":
                            info.paths += 1
                        elif kind == "FILENAME":
                            info.filenames += 1
                        elif kind == "WEB":
                            info.web_links += 1
            prof = self.match_profile(headers)
            if prof is not None:
                info.guess, info.profile = prof[1], prof[0]
                guess = info.guess
            if "photo" not in guess:
                # A column made of file names / paths is the photo column even without a "Photo" title.
                for i in range(len(headers)):
                    col = [r[i] for r in samples if i < len(r) and r[i]]
                    if col and all(_classify(v) in ("PATH", "FILENAME") for v in col) and i not in guess.values():
                        guess["photo"] = i
                        break
            sheets.append(info)
        wb.close()
        return WorkbookInfo(path, sheets)

    def read_sheet(self, path: str | Path, sheet: str, header_row: int) -> SheetInfo:
        """Column titles, sample rows and the guessed mapping when the person picks another header row."""
        self.require(P.PRODUCTS_EDIT)
        path = self._check_file(path)
        wb = self._load(path)
        try:
            ws = wb[sheet]
            headers = self._headers(ws, header_row)
            samples = self._samples(ws, header_row, len(headers))
            return SheetInfo(ws.title, max(ws.max_row - header_row, 0), header_row, headers, samples,
                             guess_mapping(headers))
        finally:
            wb.close()

    # -- preview and validation ----------------------------------------------------------------
    def prepare(self, path: str | Path, options: ImportOptions, previous: ImportPlan | None = None) -> ImportPlan:
        """Read the rows with the chosen mapping and check every one against the database.
        Nothing is changed. ``previous``: keep the person's choices (row actions, how suppliers and
        projects are resolved) when the preview is checked again."""
        self.require(P.PRODUCTS_EDIT)
        path = self._check_file(path)
        # A "Name / Description" column alone is enough: the name is taken from it.
        missing = [FIELDS[f][0].rstrip(" *") for f in REQUIRED if f not in options.mapping
                   and not (f == "name" and "description" in options.mapping)]
        if missing:
            raise ValidationError("Choose the column for: " + ", ".join(missing) + ".")
        if options.project_id and "required_qty" not in options.mapping:
            raise ValidationError("Choose the column with the Required Quantity for the project.")
        if options.project_id or "project" in options.mapping:
            self.require(P.SITES_EDIT)
        wb = self._load(path)
        if options.sheet not in wb.sheetnames:
            raise ValidationError(f"The workbook has no worksheet '{options.sheet}'.")
        ws = wb[options.sheet]
        floating = floating_pictures(ws)
        in_cell = in_cell_pictures(path).get(options.sheet, {})
        index = PhotoFolder(options.photo_folder)
        rows: list[ImportRow] = []
        photo_col = options.mapping.get("photo")
        for cells in ws.iter_rows(min_row=options.header_row + 1):
            if not cells:
                continue
            excel_row = cells[0].row
            values = {f: cell_text(cells[c].value) if c < len(cells) else "" for f, c in options.mapping.items()}
            if not any(values.values()) and excel_row not in floating:
                continue  # blank line
            if values.get("internal_code", "").strip().upper() in ("TOTAL", "TOTALS", "GRAND TOTAL") \
                    and not values.get("name"):
                continue  # the totals line under an exported list
            r = ImportRow(excel_row, values)
            r.photo = self._photo_for(r, cells, photo_col, floating.get(excel_row, []), in_cell, path, index)
            r.photo_missing = r.photo is None and any(w.startswith("photo") for w in r.warnings)
            rows.append(r)
        wb.close()
        plan = ImportPlan(path, options, rows)
        if previous is not None:
            for key, c in previous.suppliers.items():
                plan.suppliers[key] = SupplierChoice(c.text, 0, c.action, c.match_id, c.match_name, c.suggested)
            for key, c in previous.projects.items():
                plan.projects[key] = ProjectChoice(c.text, 0, c.action, c.match_id, c.match_label, c.closed)
        self._validate(plan)
        if previous is not None:
            before = {r.row: r for r in previous.rows}
            for r in plan.rows:
                old = before.get(r.row)
                if old is not None and old.code == r.code and old.action in plan.choices(r) \
                        and old.match_kind == r.match_kind:
                    r.action = old.action
        return plan

    def compare_project(self, plan: ImportPlan) -> list[CompareLine]:
        """Compare the project's current material list with the prepared file (nothing is changed).

        ADDED - in the file, not on the list; QUANTITY CHANGED - Required differs; DETAIL CHANGED -
        the material itself would be updated; UNCHANGED; REMOVED - on the list, not in the file at all
        (a row that is in the file but skipped or wrong is never reported as removed)."""
        site_id = plan.options.project_id
        if not site_id:
            raise ValidationError("Compare with Excel needs the project to import into.")
        self.require(P.SITES_VIEW)
        with self.db.read_session() as s:
            current = {pm.product_id: (pm.required_qty, p) for pm, p in s.execute(
                select(ProjectMaterial, Product).join(Product, Product.id == ProjectMaterial.product_id)
                .where(ProjectMaterial.site_id == site_id).order_by(Product.internal_code))}
            in_file = {r.existing_id for r in plan.rows if r.existing_id}
            codes = {code_key(r.code) for r in plan.rows if r.code}
            out: list[CompareLine] = []
            for r in plan.rows:
                if not plan.links_project(r, respect_excluded=False):
                    continue
                pid = None if r.action == CREATE else r.existing_id
                name = r.values.get("name", "")
                if pid in current:
                    cur, p = current[pid]
                    if r.action == UPDATE and r.changes:
                        change, details = DETAIL_CHANGED, r.changes_text
                    elif cur != r.required:
                        change, details = QTY_CHANGED, ""
                    else:
                        change, details = UNCHANGED, ""
                    out.append(CompareLine(change, p.internal_code, p.name, cur, r.required, r.row, pid, details))
                else:
                    details = {CREATE: "new material", LINK: f"matched to {r.existing_label}"}.get(r.action, "")
                    out.append(CompareLine(ADDED, r.code, name, None, r.required, r.row, pid, details))
            for pid, (cur, p) in current.items():
                if pid not in in_file and code_key(p.internal_code) not in codes:
                    out.append(CompareLine(REMOVED, p.internal_code, p.name, cur, None, None, pid,
                                           "not in the revised file"))
        order = list(COMPARE_LABELS)
        return sorted(out, key=lambda c: (order.index(c.change), c.code))

    def _photo_for(self, r: ImportRow, cells, photo_col, floating, in_cell, workbook: Path,
                   index: PhotoFolder) -> PhotoSource | None:
        # 1. a picture placed in the photo cell, or floating over that row
        if photo_col is not None and (r.row, photo_col) in in_cell:
            data, ext = in_cell[(r.row, photo_col)]
            return self._embedded(r, "IN_CELL", "picture in cell", data, ext)
        row_cells = {(row, col): v for (row, col), v in in_cell.items() if row == r.row}
        if row_cells:
            data, ext = next(iter(row_cells.values()))
            return self._embedded(r, "IN_CELL", "picture in cell", data, ext)
        if floating:
            pick = next((p for p in floating if p[0] == photo_col), floating[0])
            if len(floating) > 1:
                r.warnings.append(f"{len(floating)} pictures on this row - the one in the photo column (or the "
                                  f"first) is used")
            return self._embedded(r, "EMBEDDED", "embedded picture", pick[1], pick[2])
        if photo_col is None or photo_col >= len(cells):
            return None
        cell = cells[photo_col]
        link = cell.hyperlink.target if cell.hyperlink is not None and cell.hyperlink.target else None
        text = link or (cell.value if isinstance(cell.value, str) else "")
        kind = _classify(text or "")
        if not text:
            return None
        if kind == "WEB":
            r.warnings.append("photo is a web address - not downloaded (save the pictures to a folder and use "
                              "file names)")
            return None
        if kind is None:
            r.warnings.append(f"photo '{text}' is not a picture file (jpg / png / gif / bmp / webp)")
            return None
        found = index.find(text, workbook.parent)
        if found is None:
            r.warnings.append(f"photo file not found: {text}")
            return None
        try:
            AttachmentService(self.ctx).validate_file(found)  # damaged / renamed / too big: say so now
        except ValidationError as exc:
            r.warnings.append(f"photo not used: {exc.message}")
            return None
        return PhotoSource("HYPERLINK" if link else kind, text, path=found, ext=found.suffix.lower())

    @staticmethod
    def _embedded(r: ImportRow, kind: str, label: str, data: bytes, ext: str) -> PhotoSource | None:
        ext = ".jpg" if ext == ".jpeg" else ext
        if ext not in PHOTO_EXT:
            r.warnings.append(f"embedded picture is a {ext[1:].upper()} drawing, not a photo - not imported")
            return None
        return PhotoSource(kind, label, data=data, ext=ext)

    def _validate(self, plan: ImportPlan) -> None:
        rows, o = plan.rows, plan.options
        with self.db.read_session() as s:
            units = {u.code.upper(): u for u in s.scalars(select(Unit))}
            units.update({u.name.upper(): u for u in units.values()})
            if o.default_unit_id is None and "unit" not in o.mapping:
                raise ValidationError("Choose the unit of measure for products without a Unit column.")
            default_unit = s.get(Unit, o.default_unit_id) if o.default_unit_id else None
            existing = {p.internal_code.upper(): p for p in s.scalars(select(Product)).unique()}
            barcodes = {p.barcode.upper(): p for p in existing.values() if p.barcode}
            suppliers = {x.name.upper(): x for x in s.scalars(select(Supplier))}
            suppliers.update({x.code.upper(): x for x in suppliers.values()})
            by_key: dict[str, list[Supplier]] = {}
            for x in {x.id: x for x in suppliers.values()}.values():
                by_key.setdefault(supplier_key(x.name), []).append(x)
            by_product_no: dict[str, list[Product]] = {}
            by_code_key: dict[str, list[Product]] = {}
            for p in existing.values():
                if p.supplier_code:
                    by_product_no.setdefault(p.supplier_code.upper(), []).append(p)
                by_code_key.setdefault(code_key(p.internal_code), []).append(p)
            target = None
            if o.project_id:
                target = s.get(Site, o.project_id)
                if target is None:
                    raise ValidationError("The project to import into no longer exists.")
                if target.status in CLOSED_SITE_STATUSES:
                    raise ValidationError(f"Project {target.project_number} is "
                                          f"{target.status.lower().replace('_', ' ')}: its material list can no "
                                          f"longer be changed. Reopen it first.")
                plan.project_label = f"{target.project_number} - {target.name}"
            sites = list(s.scalars(select(Site))) if "project" in o.mapping and not o.project_id else []
            site_by_number = {x.project_number.upper(): x for x in sites}
            site_by_name: dict[str, list[Site]] = {}
            for x in sites:
                site_by_name.setdefault(x.name.upper(), []).append(x)
            categories = {n.upper() for c in s.scalars(select(Category)).unique() for n in (c.name, c.full_name)}
            bins = {}
            for b in s.scalars(select(Bin).where(Bin.is_active.is_(True))):
                bins.setdefault(b.code.upper(), []).append(b)
                bins.setdefault(b.full_code.upper(), []).append(b)
                if b.barcode:
                    bins.setdefault(b.barcode.upper(), []).append(b)
            by_supplier_code = {(p.supplier_id, (p.supplier_code or "").upper()): p for p in existing.values()
                                if p.supplier_code}
            seen: dict[str, ImportRow] = {}
            seen_barcodes: dict[str, ImportRow] = {}
            for r in rows:
                v = r.values
                code = normalise_code(v.get("internal_code"))
                v["internal_code"] = code
                if not v.get("name") and v.get("description"):
                    v["name"] = v["description"][:200]
                    if "description" in o.mapping and o.mapping.get("name") is None:
                        v["description"] = "" if len(v["description"]) <= 200 else v["description"]
                if not code:
                    r.errors.append("no code")
                elif len(code) > 50:
                    r.errors.append("code longer than 50 characters")
                if not v.get("name"):
                    r.errors.append("no name")
                elif len(v["name"]) > 200:
                    r.warnings.append("name cut to 200 characters")
                    v["name"] = v["name"][:200]
                # unit
                unit_text = (v.get("unit") or "").upper()
                unit = units.get(unit_text) if unit_text else default_unit
                if unit is None:
                    r.errors.append(f"unknown unit '{v.get('unit')}'" if unit_text else "no unit")
                else:
                    v["unit_id"] = str(unit.id)
                # numbers
                for f in (*QTY_FIELDS, "unit_cost"):
                    raw = (v.get(f) or "").replace("£", "").replace(",", "").strip()
                    if not raw:
                        v[f] = ""
                        continue
                    try:
                        num = Decimal(raw)
                    except InvalidOperation:
                        r.errors.append(f"{FIELDS[f][0]}: '{v.get(f)}' is not a number")
                        continue
                    if num < ZERO:
                        r.errors.append(f"{FIELDS[f][0]} cannot be negative")
                    elif f in QTY_FIELDS and unit is not None and not unit.allow_decimal and num != num.to_integral():
                        r.errors.append(f"{FIELDS[f][0]}: {unit.code} must be a whole number")
                    v[f] = str(num)
                # supplier / category / bin
                sup_text = (v.get("supplier") or "").strip()
                sup = suppliers.get(sup_text.upper()) if sup_text else None
                if sup_text and sup is None:
                    key = supplier_key(sup_text)
                    choice = plan.suppliers.get(key)
                    if choice is None:
                        similar = [x for x in by_key.get(key, []) if x.is_active] or by_key.get(key, [])
                        choice = SupplierChoice(sup_text, action=CREATE if o.create_suppliers else IGNORE)
                        if len(similar) == 1:
                            choice.action, choice.match_id, choice.match_name = MATCH, similar[0].id, similar[0].name
                            choice.suggested = True
                        plan.suppliers[key] = choice
                    choice.rows += 1
                    if choice.action == MATCH:
                        sup = s.get(Supplier, choice.match_id)
                        r.warnings.append(f"supplier '{sup_text}' matched to {choice.match_name}")
                    elif choice.action == CREATE:
                        r.warnings.append(f"new supplier '{sup_text}' will be created")
                    else:
                        r.warnings.append(f"supplier '{sup_text}' not in the WMS - left empty")
                elif sup is not None and not sup.is_active:
                    r.warnings.append(f"supplier {sup.name} is archived")
                cat_text = (v.get("category") or "").strip()
                if cat_text and cat_text.upper() not in categories:
                    if o.create_categories:
                        r.warnings.append(f"new category '{cat_text}' will be created")
                    else:
                        r.errors.append(f"unknown category '{cat_text}'")
                bin_text = (v.get("default_bin") or "").strip().upper()
                if bin_text:
                    found = bins.get(bin_text, [])
                    if len({b.id for b in found}) != 1:
                        r.warnings.append(f"bin '{v['default_bin']}' {'is ambiguous' if found else 'not found'} "
                                          f"- left empty")
                        v["default_bin"] = ""
                    else:
                        v["default_bin_id"] = str(found[0].id)
                # duplicates: inside the file
                if code:
                    first = seen.get(code)
                    if first is not None:
                        r.warnings.append(f"duplicate: code {code} is also on row {first.row} - this row is "
                                          f"not imported")
                        r.duplicate_of = first.row
                        r.action = SKIP
                    else:
                        seen[code] = r
                bc = (v.get("barcode") or "").upper()
                if bc:
                    other = seen_barcodes.get(bc)
                    if other is not None:
                        r.errors.append(f"barcode {v['barcode']} is also on row {other.row}")
                    seen_barcodes[bc] = r
                # duplicates: already in the WMS
                match = existing.get(code) if code else None
                kind = "CODE" if match is not None else None
                sc = (v.get("supplier_code") or "").upper()
                if match is None and sc:
                    # no such code: is it the same product under another code? Product number (and the
                    # supplier, when both sides have one) - never the description alone (spec section 31)
                    twins = [p for p in by_product_no.get(sc, [])
                             if sup is None or p.supplier_id in (None, sup.id)]
                    if len(twins) == 1:
                        match, kind = twins[0], "PRODUCT_NO"
                if match is None and code:
                    # 'H3311 ST10' in the file, 'H3311-ST10' in the WMS: probably the same material, but
                    # only a suggestion - never merged without the person's choice (spec 31)
                    similar = [p for p in by_code_key.get(code_key(code), []) if p.internal_code.upper() != code]
                    if len(similar) == 1:
                        match, kind = similar[0], "SIMILAR_CODE"
                if match is not None:
                    r.existing_id = match.id
                    r.match_kind = kind
                    r.existing_label = f"{match.internal_code} - {match.name}"
                    r.changes = self._changes(match, v, o, sup)
                    if kind == "PRODUCT_NO":
                        r.warnings.append(f"same product number as {match.internal_code} - possible duplicate "
                                          f"(choose {'Match to use it, ' if plan.has_project else ''}Update to "
                                          f"change {match.internal_code}, or New to add it anyway)")
                    elif kind == "SIMILAR_CODE":
                        r.warnings.append(f"possible existing material {match.internal_code} - {match.name} "
                                          f"(choose {'Match to use it, ' if plan.has_project else ''}Update to "
                                          f"change it, or New to create a separate material)")
                    if not match.is_active:
                        r.action = SKIP
                        r.warnings.append(f"{match.internal_code} exists but was removed from active use - "
                                          f"restore it first to update it")
                    elif kind in UNCERTAIN:
                        r.action = SKIP
                    else:
                        r.action = o.on_existing if o.on_existing in (SKIP, UPDATE) else SKIP
                        if unit is not None and match.unit_id != unit.id and r.action == UPDATE:
                            r.warnings.append(f"unit stays {match.unit.code} (not changed by import)")
                elif r.ok and r.duplicate_of is None:
                    r.action = CREATE
                if bc and bc in barcodes and barcodes[bc].id != r.existing_id:
                    r.errors.append(f"barcode {v['barcode']} already belongs to {barcodes[bc].internal_code}")
                if bc and bc in existing and existing[bc].id != r.existing_id:
                    r.errors.append(f"barcode {v['barcode']} is the code of product {bc}")
                if code and code in barcodes and barcodes[code].id != r.existing_id:
                    r.errors.append(f"{code} is the barcode of product {barcodes[code].internal_code}")
                if sc and sup is not None and r.existing_id is None:
                    twin = by_supplier_code.get((sup.id, sc))
                    if twin is not None:
                        r.warnings.append(f"same supplier code as {twin.internal_code} - possible duplicate")
                self._project_part(plan, r, unit, target, site_by_number, site_by_name)
                if not r.ok or r.duplicate_of is not None:
                    r.action = SKIP

    @staticmethod
    def _changes(p: Product, v: dict, o: ImportOptions, sup: Supplier | None) -> list[tuple[str, str, str]]:
        """What an Update would change: (field, now in the WMS, from Excel). Only mapped, filled cells."""
        out = []

        def num(x) -> str:
            return "" if x is None else format(Decimal(x).normalize(), "f")

        for f in ("name", "description", "supplier_code", "barcode", "brand", "notes"):
            new = (v.get(f) or "").strip()
            if f in o.mapping and new and new != (getattr(p, f) or ""):
                out.append((FIELDS[f][0].rstrip(" *"), getattr(p, f) or "", new))
        for f in (*QTY_FIELDS, "unit_cost"):
            new = v.get(f) or ""
            if f in o.mapping and new:
                try:
                    if Decimal(new) != (getattr(p, f) if getattr(p, f) is not None else None):
                        out.append((FIELDS[f][0], num(getattr(p, f)), num(new)))
                except InvalidOperation:
                    pass
        if sup is not None and sup.id != p.supplier_id:
            out.append(("Supplier", p.supplier.name if p.supplier else "", sup.name))
        elif (v.get("supplier") or "").strip() and sup is None:
            out.append(("Supplier", p.supplier.name if p.supplier else "", v["supplier"].strip()))
        cat = (v.get("category") or "").strip()
        if cat and (p.category is None or cat.upper() not in (p.category.name.upper(), p.category.full_name.upper())):
            out.append(("Category", p.category.full_name if p.category else "", cat))
        if v.get("default_bin_id") and int(v["default_bin_id"]) != p.default_bin_id:
            out.append(("Default bin", p.default_bin.full_code if p.default_bin else "", v.get("default_bin", "")))
        return out

    @staticmethod
    def _project_part(plan: ImportPlan, r: ImportRow, unit, target, by_number, by_name) -> None:
        """Required quantity and which project the row's material goes on (spec sections 34-35)."""
        o, v = plan.options, r.values
        raw = (v.get("required_qty") or "").replace(",", "").strip()
        wants_project = bool(o.project_id) or bool((v.get("project") or "").strip())
        if raw:
            try:
                q = Decimal(raw)
            except InvalidOperation:
                r.errors.append(f"Required Quantity '{v.get('required_qty')}' is not a number")
                q = None
            if q is not None and q < 0:
                r.errors.append("Required Quantity cannot be negative")
            elif q is not None and unit is not None and not unit.allow_decimal and q != q.to_integral():
                r.errors.append(f"Required Quantity: {unit.code} must be a whole number")
            elif q is not None:
                r.required = q
        elif wants_project and "required_qty" in o.mapping:
            r.warnings.append("no Required Quantity - not added to the project list")
        if o.project_id or not (v.get("project") or "").strip():
            if raw and not o.project_id and "project" not in o.mapping:
                r.warnings.append("Required Quantity ignored: no project (map the Project column or use "
                                  "Import Project Material Requirements)")
            return
        text = v["project"].strip()
        key = text.upper()
        r.project_key = key
        choice = plan.projects.get(key)
        if choice is None:
            choice = ProjectChoice(text)
            site = by_number.get(key)
            if site is None and len(by_name.get(key, [])) == 1:
                site = by_name[key][0]
            if site is not None:
                choice.match_id, choice.match_label = site.id, f"{site.project_number} - {site.name}"
                choice.closed = site.status in CLOSED_SITE_STATUSES
                choice.action = IGNORE if choice.closed else MATCH
            plan.projects[key] = choice
        choice.rows += 1
        if r.required is None:
            return
        if choice.action == MATCH:
            r.warnings.append(f"goes on project {choice.match_label}: required {format(r.required.normalize(), 'f')}")
        elif choice.closed:
            r.warnings.append(f"project {choice.match_label} is closed - not added to its list")
        elif choice.action == IGNORE:
            r.warnings.append(f"project '{text}' not in the WMS - material imported without it")

    # -- the import ---------------------------------------------------------------------------
    def run(self, plan: ImportPlan, progress=None) -> ImportResult:
        """Import the rows marked New / Update (and their project list lines). Backup first; then
        everything in one transaction. Stock is never changed (spec section 33).

        ``progress(done, total)`` is called as rows are written; returning False cancels, and the
        transaction is rolled back - nothing is imported."""
        self.require(P.PRODUCTS_EDIT)
        rows = plan.to_import
        if not rows and not (plan.remove_ids and plan.options.project_id):
            raise ValidationError("Nothing to import: every row is skipped or has errors.")
        for r in rows:
            if r.action == UPDATE and r.existing_id is None:
                raise ValidationError(f"Row {r.row} ({r.code}): only a material already in the WMS can be "
                                      f"updated. Check the preview again.")
            if r.action == CREATE and r.match_kind == "CODE":
                raise ValidationError(f"Row {r.row} ({r.code}): the code is already in the WMS - choose Update "
                                      f"or Skip. Check the preview again.")
        links = [r for r in rows if plan.links_project(r)]
        if links or plan.remove_ids:
            self.require(P.SITES_EDIT)
        with_photos = [r for r in rows if r.photo is not None and r.action in (CREATE, UPDATE)]
        if with_photos:
            self.require(P.ATTACHMENTS_ADD)
        from wms.services.backup import BackupService
        try:
            backup = BackupService(self.ctx).create_safety(
                f"Pre-import backup before Excel import of {plan.source.name} ({len(rows)} products)")
        except WMSError:
            raise
        except Exception as exc:  # noqa: BLE001 - no backup, no import
            raise ValidationError(f"The pre-import backup failed, so nothing was imported: {exc}") from exc
        att = AttachmentService(self.ctx)
        staged: dict[int, StagedFile] = {}
        photo_problems: list[str] = []
        with tempfile.TemporaryDirectory() as tmp:
            for r in with_photos:
                src = r.photo.path
                if src is None:
                    src = Path(tmp) / f"row{r.row}_{r.code or 'photo'}{r.photo.ext}"
                    src = src.with_name(re.sub(r"[^A-Za-z0-9._-]+", "_", src.name))
                    src.write_bytes(r.photo.data)
                try:
                    staged[r.row] = att.stage(src)
                except WMSError as exc:
                    photo_problems.append(f"Row {r.row} ({r.code}): {exc.message}")
        products = ProductService(self.ctx)
        sup_new: dict[str, Supplier] = {}
        cat_new: dict[str, Category] = {}
        site_new: dict[str, Site] = {}
        created = updated = photos = 0
        project_items: dict[int, list[tuple[int, Decimal]]] = {}
        total = len(rows)
        try:
            with self.db.session() as s:
                suppliers = {x.name.upper(): x for x in s.scalars(select(Supplier))}
                suppliers.update({x.code.upper(): x for x in list(suppliers.values())})
                categories = {n.upper(): c for c in s.scalars(select(Category)).unique() for n in (c.name, c.full_name)}
                for n, r in enumerate(rows, start=1):
                    if progress is not None and n % 25 == 1 and progress(n - 1, total) is False:
                        raise ImportCancelled()
                    v = r.values
                    p_id = None if r.action == CREATE else r.existing_id
                    if r.action in (CREATE, UPDATE):
                        try:
                            data = self._product_data(s, r, plan, suppliers, sup_new, categories, cat_new)
                            p = products.save_in(s, r.existing_id if r.action == UPDATE else None, data)
                        except WMSError as exc:
                            raise ValidationError(f"Row {r.row} ({v.get('internal_code')}): {exc.message} "
                                                  f"- nothing was imported.") from exc
                        p_id = p.id
                        if r.action == UPDATE:
                            updated += 1
                        else:
                            created += 1
                    if plan.links_project(r):
                        site_id = self._project_for(s, plan, r, site_new)
                        project_items.setdefault(site_id, []).append((p_id, r.required))
                    f = staged.get(r.row)
                    if f is None or r.action not in (CREATE, UPDATE):
                        continue
                    if f is not None:
                        same = s.scalar(select(func.count(Attachment.id)).where(
                            Attachment.entity_type == EntityType.PRODUCT.value, Attachment.entity_id == p.id,
                            Attachment.sha256 == f.sha256, Attachment.is_deleted.is_(False)))
                        if same:
                            att.discard([f])
                            staged.pop(r.row)
                            continue
                        att.attach_staged(s, EntityType.PRODUCT.value, p.id, f, caption="Imported from Excel")
                        photos += 1
                from wms.services.materials import upsert_project_materials
                lines = 0
                for site_id, items in project_items.items():
                    try:
                        upsert_project_materials(s, self.ctx, site_id, items, replace=True,
                                                 source=f"Excel import {plan.source.name}")
                    except WMSError as exc:
                        raise ValidationError(f"{exc.message} Nothing was imported.") from exc
                    lines += len({pid for pid, _ in items})
                removed = 0
                if plan.remove_ids and plan.options.project_id:
                    # Compare with Excel: lines the person confirmed to take off (spec 33). Only the
                    # Required quantity goes; reservations and issued history stay.
                    res = s.execute(delete(ProjectMaterial).where(
                        ProjectMaterial.site_id == plan.options.project_id,
                        ProjectMaterial.product_id.in_(plan.remove_ids)))
                    removed = res.rowcount or 0
                    site = s.get(Site, plan.options.project_id)
                    audit.record(s, self.ctx, "PROJECT_MATERIALS_REMOVED", EntityType.SITE.value, site.id,
                                 site.project_number, f"{removed} material(s) taken off the project list "
                                 f"(not in revised Excel {plan.source.name})",
                                 old={"products": sorted(plan.remove_ids)[:500]})
                if progress is not None and progress(total, total) is False:
                    raise ImportCancelled()
                counts = plan.counts()
                audit.record(s, self.ctx, "EXCEL_IMPORT", "PRODUCT", None, plan.source.name,
                             f"Excel import of {plan.source.name} / {plan.options.sheet}: {created} created, "
                             f"{updated} updated, {counts['skip']} skipped, {counts['errors']} rows with errors, "
                             f"{photos} photo(s). Backup {backup.name}.",
                             new={"created": created, "updated": updated, "skipped": counts["skip"],
                                  "errors": counts["errors"], "photos": photos, "backup": backup.name,
                                  "suppliers_created": sorted(sup_new), "categories_created": sorted(cat_new),
                                  "project_lines": lines, "projects_created": sorted(site_new),
                                  "project_lines_removed": removed})
        except BaseException:
            att.discard(list(staged.values()))
            raise
        c = plan.counts()
        return ImportResult(created, updated, c["skip"], c["errors"], photos,
                            sorted(x.name for x in sup_new.values()), sorted(c_.name for c_ in cat_new.values()),
                            backup, photo_problems, duplicates=c["duplicate"],
                            photos_missing=c["photos_missing"] + len(photo_problems), project_lines=lines,
                            projects_created=sorted(f"{x.project_number} - {x.name}" for x in site_new.values()),
                            total_rows=c["rows"], project_lines_removed=removed)

    def _project_for(self, s, plan: ImportPlan, r: ImportRow, site_new: dict[str, Site]) -> int:
        if plan.options.project_id:
            return plan.options.project_id
        c = plan.projects[r.project_key]
        if c.action == MATCH:
            return c.match_id
        site = site_new.get(r.project_key)
        if site is None:
            number = normalise_code(c.text)[:50]
            if s.scalar(select(Site.id).where(func.upper(Site.project_number) == number.upper())):
                raise ValidationError(f"Project {number} already exists - choose 'Match' for it. Nothing was "
                                      f"imported.")
            site = Site(project_number=number, name=c.text[:200], status=SiteStatus.ACTIVE.value,
                        notes=f"Created by Excel import of {plan.source.name}")
            s.add(site)
            s.flush()
            audit.record(s, self.ctx, "SITE_SAVED", EntityType.SITE.value, site.id, site.project_number,
                         "Created by Excel import")
            site_new[r.project_key] = site
        return site.id

    def _product_data(self, s, r: ImportRow, plan: ImportPlan, suppliers, sup_new, categories, cat_new) -> dict:
        v = r.values
        o = plan.options
        mapped = set(o.mapping)
        if r.action == UPDATE:
            p = s.get(Product, r.existing_id)
            # Start from the product as it is; only columns that were mapped AND filled change it.
            data = {f: getattr(p, f) for f in ("item_type", "internal_code", "name", "supplier_code", "barcode",
                                                "description", "brand", "notes", "unit_id", "category_id",
                                                "supplier_id", "default_bin_id", "min_stock", "reorder_qty",
                                                "reorder_level", "unit_cost", "is_active")}
            if p.finished_info is not None:
                fi = p.finished_info
                data.update(production_date=fi.production_date, site_id=fi.site_id, customer=fi.customer,
                            status=fi.status)
        else:
            data = {"item_type": o.item_type, "internal_code": v["internal_code"], "unit_id": int(v["unit_id"]),
                    "is_active": True}
        for f in ("name", "description", "supplier_code", "barcode", "brand", "notes", *QTY_FIELDS, "unit_cost"):
            if f in mapped and v.get(f):
                data[f] = v[f]
            elif r.action == CREATE:
                data.setdefault(f, None)
        if v.get("default_bin_id"):
            data["default_bin_id"] = int(v["default_bin_id"])
        sup_text = (v.get("supplier") or "").strip()
        if sup_text:
            sup = suppliers.get(sup_text.upper())
            if sup is None:
                key = supplier_key(sup_text)
                choice = plan.suppliers.get(key) or SupplierChoice(sup_text, action=IGNORE)
                if choice.action == MATCH and choice.match_id:
                    sup = s.get(Supplier, choice.match_id)
                elif choice.action == CREATE:
                    sup = sup_new.get(key)  # "EGGER" and "Egger" in one file: one new supplier
                    if sup is None:
                        sup = self._new_supplier(s, choice.text, suppliers)
                        sup_new[key] = sup
            if sup is not None:
                data["supplier_id"] = sup.id
        cat_text = (v.get("category") or "").strip()
        if cat_text:
            cat = categories.get(cat_text.upper())
            if cat is None:
                cat = Category(name=cat_text[:100], is_active=True)
                s.add(cat)
                s.flush()
                audit.record(s, self.ctx, "CATEGORY_SAVED", "CATEGORY", cat.id, cat.name, "Created by Excel import")
                categories[cat_text.upper()] = cat
                cat_new[cat_text.upper()] = cat
            data["category_id"] = cat.id
        return data

    def _new_supplier(self, s, name: str, suppliers: dict) -> Supplier:
        base = re.sub(r"[^A-Z0-9]", "", name.upper())[:10] or "SUP"
        code, n = base, 1
        while code in suppliers or s.scalar(select(Supplier.id).where(func.upper(Supplier.code) == code)):
            n += 1
            code = f"{base[:8]}{n}"
        sup = Supplier(code=code, name=name[:200], is_active=True)
        s.add(sup)
        s.flush()
        audit.record(s, self.ctx, "SUPPLIER_CREATED", "SUPPLIER", sup.id, name, "Created by Excel import")
        suppliers[name.upper()] = sup
        suppliers[code] = sup
        return sup

    # -- mapping profiles (spec section 29) ---------------------------------------------------
    def profiles(self) -> list[tuple[int, str, dict[str, str]]]:
        """(id, name, field -> column title) for the person's saved column mappings."""
        out = []
        with self.db.read_session() as s:
            for row in s.scalars(select(SavedSearch).where(SavedSearch.user_id == self.user.id,
                                                           SavedSearch.criteria.startswith(PROFILE_PREFIX))
                                 .order_by(func.lower(SavedSearch.name))):
                try:
                    data = json.loads(row.criteria)["import_profile"]
                except (ValueError, KeyError, TypeError):
                    continue
                out.append((row.id, data.get("name") or row.name, dict(data.get("mapping") or {})))
        return out

    def save_profile(self, name: str, headers: list[str], mapping: dict[str, int]) -> int:
        """Remember a mapping by column TITLE, so it fits the next file with the same columns even if
        they are in another order."""
        self.require(P.PRODUCTS_EDIT)
        from wms.services.base import clean_str
        name = clean_str(name, "name", required=True, max_len=60)
        by_title = {f: headers[c] for f, c in mapping.items() if 0 <= c < len(headers)}
        if not by_title:
            raise ValidationError("Map at least one column before saving the profile.")
        text = json.dumps({"import_profile": {"name": name, "mapping": by_title}}, sort_keys=True)
        stored = f"Import: {name}"[:80]
        with self.db.session() as s:
            row = s.scalar(select(SavedSearch).where(SavedSearch.user_id == self.user.id,
                                                     func.lower(SavedSearch.name) == stored.lower()))
            if row is None:
                row = SavedSearch(user_id=self.user.id, name=stored, criteria=text)
                s.add(row)
            else:
                row.criteria = text
            s.flush()
            return row.id

    def delete_profile(self, profile_id: int) -> None:
        with self.db.session() as s:
            row = s.get(SavedSearch, profile_id)
            if row is not None and row.user_id == self.user.id and row.criteria.startswith(PROFILE_PREFIX):
                s.delete(row)

    def match_profile(self, headers: list[str]) -> tuple[str, dict[str, int]] | None:
        """The saved profile whose every column title is in ``headers`` (most columns wins)."""
        if self.user is None:
            return None
        normed = {_norm_header(h): i for i, h in enumerate(headers) if h}
        best = None
        for _id, name, by_title in self.profiles():
            cols = {f: normed.get(_norm_header(t)) for f, t in by_title.items()}
            if by_title and all(c is not None for c in cols.values()) and (best is None or len(cols) > len(best[1])):
                best = (name, cols)
        return best

    def apply_profile(self, profile_id: int, headers: list[str]) -> dict[str, int]:
        """Field -> column for a chosen profile (titles not in this file are left unmapped)."""
        normed = {_norm_header(h): i for i, h in enumerate(headers) if h}
        for pid, _name, by_title in self.profiles():
            if pid == profile_id:
                return {f: normed[_norm_header(t)] for f, t in by_title.items() if _norm_header(t) in normed}
        raise ValidationError("That mapping profile no longer exists.")

    # -- helpers --------------------------------------------------------------------------------
    @staticmethod
    def _check_file(path: str | Path) -> Path:
        p = Path(path)
        if not p.is_file():
            raise ValidationError(f"File not found: {p}")
        if p.suffix.lower() not in (".xlsx", ".xlsm"):
            raise ValidationError("Choose an Excel workbook (.xlsx). Older .xls files: open them in Excel and "
                                  "'Save As' .xlsx first.")
        return p

    @staticmethod
    def _load(path: Path):
        from openpyxl import load_workbook
        try:
            return load_workbook(path, data_only=True)
        except zipfile.BadZipFile:
            raise ValidationError(f"'{path.name}' is not a valid Excel workbook (or it is password protected).") \
                from None
        except PermissionError:
            raise ValidationError(f"'{path.name}' is open in another program. Close it and try again.") from None

    @staticmethod
    def _samples(ws, header_row: int, width: int, count: int = 5) -> list[list[str]]:
        """The first filled-in rows under the header, as text."""
        out = []
        for r in ws.iter_rows(min_row=header_row + 1, max_col=width, values_only=True):
            texts = [cell_text(v) for v in r]
            if any(texts):
                out.append(texts)
                if len(out) == count:
                    break
        return out

    @staticmethod
    def _headers(ws, header_row: int) -> list[str]:
        values = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        last = max((i for i, v in enumerate(values) if v not in (None, "")), default=-1)
        width = max(last + 1, min(ws.max_column, 40))
        from openpyxl.utils import get_column_letter
        return [cell_text(values[i]) if i < len(values) and values[i] not in (None, "")
                else f"Column {get_column_letter(i + 1)}" for i in range(width)]


class ImportCancelled(ValidationError):
    def __init__(self):
        super().__init__("Import cancelled - nothing was imported (the WMS is exactly as before).")


class PhotoFolder:
    """Finds a photo by its path or file name: as written, next to the workbook, or anywhere in the
    chosen picture folder (by file name, ignoring case)."""

    def __init__(self, folder: Path | None):
        self.folder = Path(folder) if folder else None
        self._by_name: dict[str, list[Path]] | None = None

    def _index(self) -> dict[str, list[Path]]:
        if self._by_name is None:
            self._by_name = {}
            if self.folder and self.folder.is_dir():
                for f in self.folder.rglob("*"):
                    if f.is_file() and f.suffix.lower() in PHOTO_EXT:
                        self._by_name.setdefault(f.name.lower(), []).append(f)
        return self._by_name

    def find(self, text: str, workbook_dir: Path) -> Path | None:
        t = text.strip()
        if t.lower().startswith("file:"):
            from urllib.parse import unquote, urlparse
            t = unquote(urlparse(t).path)
            if re.match(r"^/[A-Za-z]:", t):
                t = t[1:]
        candidate = Path(t.replace("\\", "/")) if "\\" in t and not Path(t).exists() else Path(t)
        for p in (candidate, workbook_dir / candidate):
            if p.is_file():
                return p
        name = Path(t.replace("\\", "/")).name.lower()
        hits = self._index().get(name, [])
        if len(hits) == 1:
            return hits[0]
        if not hits and self.folder is not None and (self.folder / name).is_file():
            return self.folder / name
        return hits[0] if hits else None


#: template columns (spec section 40): the important ones first, then the rest of FIELDS
TEMPLATE_COLUMNS = [("internal_code", "Code"), ("name", "Material Name"), ("supplier_code", "Product Number"),
                    ("supplier", "Supplier"), ("photo", "Photo"), ("category", "Category"), ("brand", "Brand"),
                    ("project", "Project"), ("required_qty", "Required Quantity"), ("description", "Description"),
                    ("barcode", "Barcode"), ("unit", "Unit"), ("min_stock", "Minimum stock"),
                    ("reorder_level", "Reorder level"), ("reorder_qty", "Reorder quantity"),
                    ("unit_cost", "Unit cost (£)"), ("default_bin", "Default bin"), ("notes", "Notes")]

TEMPLATE_HELP = [
    ("Code", "Required. Your internal material code, e.g. H3311-18. Rows are matched to existing materials by "
             "this code first."),
    ("Material Name", "Required. What the material is called, e.g. Egger H3311 Halifax Oak 18mm."),
    ("Product Number", "The supplier's own product / decor number. Used to spot the same material under another "
                       "code."),
    ("Supplier", "Supplier name. 'EGGER' and 'Egger UK Ltd' are recognised as the same supplier; unknown "
                 "suppliers are listed for you to Match, Create or Ignore before anything is imported."),
    ("Photo", "A picture placed in the cell, a file name (choose the photo folder when importing), a full path "
              "or a hyperlink to the file. Photos are copied into the WMS."),
    ("Project", "Optional project number or name. The material goes on that project's material list with the "
                "Required Quantity. Unknown projects are listed for you to decide."),
    ("Required Quantity", "How much the project needs. This is NOT stock: importing never changes stock, "
                          "reservations or Goods Out."),
    ("Unit", "SHEET, M, M2, EACH, BOX... Leave empty to use the default unit chosen in the import."),
    ("Numbers", "Quantities and prices as plain numbers (12, 24.50). Prices in pounds (£)."),
]


def copy_template(dest: Path) -> Path:
    """A blank workbook with the columns the importer understands, plus a sheet explaining them."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    fill = PatternFill("solid", fgColor="1F3A5F")
    for i, (_f, title) in enumerate(TEMPLATE_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(title) + 4)
    ws.column_dimensions["B"].width = 40
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TEMPLATE_COLUMNS))}1"
    help_ws = wb.create_sheet("How to fill in")
    help_ws.append(["Column", "What to put in it"])
    for c in help_ws[1]:
        c.font = Font(bold=True)
    for row in TEMPLATE_HELP:
        help_ws.append(list(row))
    help_ws.column_dimensions["A"].width = 20
    help_ws.column_dimensions["B"].width = 110
    for r in help_ws.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(dest)
    return dest


def export_error_report(plan: ImportPlan, dest: Path) -> Path:
    """Rows with errors, duplicates or notes, with the reason - to fix in the original workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Import problems"
    ws.append([f"Import of {plan.source.name} / {plan.options.sheet}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([plan.summary()])
    ws.append([])
    ws.append(["Excel row", "Status", "Code", "Material", "Product Number", "Supplier", "Problem / note",
               "Changes"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in plan.rows:
        if r.errors or r.warnings or r.status in (DUPLICATE, ERROR):
            ws.append([r.row, r.status, r.code, r.values.get("name", ""), r.values.get("supplier_code", ""),
                       r.values.get("supplier", ""), r.problems, r.changes_text])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{max(ws.max_row, 4)}"
    for col, w in zip("ABCDEFGH", (10, 12, 14, 36, 16, 20, 70, 50), strict=True):
        ws.column_dimensions[col].width = w
    wb.save(dest)
    return dest


__all__ = ["ACTION_LABELS", "CREATE", "DUPLICATE", "ERROR", "FIELDS", "IGNORE", "MATCH", "NEW", "PROJECT_ACTIONS",
           "SKIP", "SKIPPED", "SKIP_ROW", "SUPPLIER_ACTIONS", "UPDATE", "ExcelImportService", "ImportCancelled",
           "ImportOptions", "ImportPlan", "ImportResult", "ImportRow", "ProjectChoice", "SheetInfo", "SupplierChoice",
           "WorkbookInfo", "copy_template", "export_error_report", "guess_mapping", "supplier_key"]
