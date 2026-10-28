"""Export a ReportResult to CSV, Excel (.xlsx) or PDF."""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from wms import APP_NAME
from wms.core.numbers import fmt_date, fmt_datetime, fmt_money, fmt_qty, to_local
from wms.services.reports import DATE, DATETIME, INT, MONEY, QTY, ReportResult


def _text(value, kind: str) -> str:
    if value is None:
        return ""
    if kind == QTY:
        return fmt_qty(value)
    if kind == MONEY:
        return fmt_money(value)
    if kind == DATE:
        return fmt_date(value)
    if kind == DATETIME:
        return fmt_datetime(value)
    return str(value)


def export_csv(result: ReportResult, path: Path) -> Path:
    # utf-8-sig so Excel opens £ and accented names correctly
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c.title for c in result.columns])
        for row in result.rows:
            out = []
            for v, c in zip(row, result.columns):
                if v is None:
                    out.append("")
                elif c.kind in (QTY, MONEY, INT):
                    out.append(format(Decimal(v).normalize(), "f") if isinstance(v, Decimal) else v)
                elif c.kind == DATE:
                    out.append(fmt_date(v))
                elif c.kind == DATETIME:
                    out.append(fmt_datetime(v))
                else:
                    out.append(v)
            w.writerow(out)
    return path


def export_xlsx(result: ReportResult, path: Path, company: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = result.title[:31]
    ws["A1"] = result.title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"{company + ' - ' if company else ''}Generated {datetime.now():%d/%m/%Y %H:%M}"
    ws["A3"] = result.filters_text or "No filters"
    header_row = 5
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for i, c in enumerate(result.columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=c.title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    formats = {QTY: "#,##0.###", MONEY: "£#,##0.00", DATE: "dd/mm/yyyy", DATETIME: "dd/mm/yyyy hh:mm", INT: "0"}
    for r, row in enumerate(result.rows, start=header_row + 1):
        for i, (v, c) in enumerate(zip(row, result.columns), start=1):
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, datetime):
                v = to_local(v).replace(tzinfo=None)
            cell = ws.cell(row=r, column=i, value=v)
            if c.kind in formats:
                cell.number_format = formats[c.kind]
    totals = result.totals()
    if any(t is not None for t in totals):
        tr = header_row + len(result.rows) + 1
        ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
        for i, (t, c) in enumerate(zip(totals, result.columns), start=1):
            if t is not None:
                cell = ws.cell(row=tr, column=i, value=float(t))
                cell.font = Font(bold=True)
                cell.number_format = formats.get(c.kind, "General")
    for i, c in enumerate(result.columns, start=1):
        width = max([len(c.title)] + [len(_text(row[i - 1], c.kind)) for row in result.rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(result.columns))}{header_row + len(result.rows)}"
    wb.save(path)
    return path


def export_pdf(result: ReportResult, path: Path, company: str = "") -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    page = landscape(A4) if len(result.columns) > 6 else A4
    doc = SimpleDocTemplate(str(path), pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=result.title, author=APP_NAME)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    head_style = cell_style.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(result.title, styles["Title"]),
             Paragraph(f"{company + ' - ' if company else ''}Generated {datetime.now():%d/%m/%Y %H:%M}"
                       + (f" - {result.filters_text}" if result.filters_text else ""), styles["Normal"]),
             Spacer(1, 4 * mm)]

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    data = [[Paragraph(esc(c.title), head_style) for c in result.columns]]
    for row in result.rows:
        data.append([Paragraph(esc(_text(v, c.kind)), cell_style) for v, c in zip(row, result.columns)])
    totals = result.totals()
    if any(t is not None for t in totals) and result.rows:
        data.append([Paragraph("<b>Total</b>" if i == 0 else (f"<b>{esc(_text(t, c.kind))}</b>" if t is not None else ""),
                               cell_style) for i, (t, c) in enumerate(zip(totals, result.columns))])
    avail = page[0] - 24 * mm
    weights = []
    for i, c in enumerate(result.columns):
        longest = max([len(c.title)] + [len(_text(r[i], c.kind)) for r in result.rows[:300]])
        weights.append(min(max(longest, 5), 40))
    total_w = sum(weights) or 1
    widths = [avail * w / total_w for w in weights]
    table = LongTable(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8CDD3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
    ]
    table.setStyle(TableStyle(style))
    story.append(table)
    if not result.rows:
        story.append(Paragraph("No records match the selected filters.", styles["Normal"]))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(page[0] - 12 * mm, 8 * mm, f"Page {d.page}")
        canvas.drawString(12 * mm, 8 * mm, f"{APP_NAME} - {result.title}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def export(result: ReportResult, path: Path, company: str = "") -> Path:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return export_csv(result, path)
    if ext == ".xlsx":
        return export_xlsx(result, path, company)
    if ext == ".pdf":
        return export_pdf(result, path, company)
    raise ValueError(f"Unsupported export format: {ext}")


__all__ = ["export", "export_csv", "export_xlsx", "export_pdf"]
