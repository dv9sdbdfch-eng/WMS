"""Excel export of the Materials list (spec sections 20-23).

Three scopes, one source of rows:

* Current View  - ``MaterialsService.iter_rows(filter)``: the same statement the table runs, so
  the file holds exactly the rows on screen, in the same order;
* Selected      - ``MaterialsService.rows_for(ids)``: the ticked rows;
* All Materials - ``iter_rows`` with the "everything, active and archived" filter.

The workbook is written in openpyxl's streaming (write-only) mode: memory stays flat however many
materials there are. It has a title block, a bold header row, AutoFilter, a frozen header,
sensible widths, real numbers for quantities and real dates.
"""

from __future__ import annotations

from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from wms.core.numbers import to_local
from wms.services.materials import MaterialFilter, MaterialRow, MaterialsService
from wms.services.reports import DATE, DATETIME, QTY, TEXT

CURRENT, SELECTED, ALL = "CURRENT", "SELECTED", "ALL"


@dataclass
class Col:
    title: str
    kind: str
    width: int
    get: Callable[[MaterialRow], object]


def _qty(v: Decimal | None):
    return None if v is None else float(v)


#: Column titles match the Excel import's column names, so an exported list can be corrected in
#: Excel and imported back (the import never changes stock, so the stock columns are ignored).
MATERIAL_COLUMNS = [
    Col("Code", TEXT, 14, lambda r: r.product.internal_code),
    Col("Material", TEXT, 36, lambda r: r.product.name),
    Col("Description", TEXT, 30, lambda r: r.product.description or ""),
    Col("Product Number", TEXT, 16, lambda r: r.product.supplier_code or ""),
    Col("Supplier", TEXT, 20, lambda r: r.product.supplier.name if r.product.supplier else ""),
    Col("Barcode", TEXT, 15, lambda r: r.product.barcode or ""),
    Col("Brand", TEXT, 14, lambda r: r.product.brand or ""),
    Col("Category", TEXT, 18, lambda r: r.product.category.full_name if r.product.category else ""),
    Col("Unit", TEXT, 7, lambda r: r.product.unit.code),
    Col("On Hand", QTY, 10, lambda r: _qty(r.on_hand)),
    Col("Reserved", QTY, 10, lambda r: _qty(r.reserved)),
    Col("Available", QTY, 10, lambda r: _qty(r.available)),
    Col("Incoming", QTY, 10, lambda r: _qty(r.incoming)),
    Col("Bins", TEXT, 24, lambda r: r.bins_detail),
    Col("Default Bin", TEXT, 14, lambda r: r.product.default_bin.full_code if r.product.default_bin else ""),
    Col("Stock Status", TEXT, 13, lambda r: r.status_label),
    Col("Last Changed", DATETIME, 16, lambda r: r.product.updated_at),
]

PROJECT_COLUMNS = [
    Col("Internal Code", TEXT, 14, lambda r: r.product.internal_code),
    Col("Material", TEXT, 36, lambda r: r.product.name),
    Col("Product Number", TEXT, 16, lambda r: r.product.supplier_code or ""),
    Col("Supplier", TEXT, 20, lambda r: r.product.supplier.name if r.product.supplier else ""),
    Col("Unit", TEXT, 7, lambda r: r.product.unit.code),
    Col("Required", QTY, 10, lambda r: _qty(r.required)),
    Col("Reserved", QTY, 10, lambda r: _qty(r.proj_reserved)),
    Col("Issued", QTY, 10, lambda r: _qty(r.issued)),
    Col("Returned", QTY, 10, lambda r: _qty(r.returned)),
    Col("Net Issued", QTY, 11, lambda r: _qty(r.net)),
    Col("Available in Stock", QTY, 12, lambda r: _qty(r.available)),
]

FORMATS = {QTY: "#,##0.###", DATE: "dd/mm/yyyy", DATETIME: "dd/mm/yyyy hh:mm"}


def write_workbook(path: Path, title: str, info: list[str], columns: list[Col],
                   chunks: Iterable[list[MaterialRow]], total: int | None = None,
                   progress: Callable[[int, int | None], bool] | None = None) -> int:
    """Stream rows into an .xlsx file. ``progress(done, total)`` returning False cancels (the
    partial file is deleted). Returns the number of rows written."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31] or "Materials")
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.width
    head = len(info) + 3  # header row: after the title, the info lines and a blank line
    ws.freeze_panes = f"A{head + 1}"
    t = WriteOnlyCell(ws, value=title)
    t.font = Font(size=14, bold=True)
    ws.append([t])
    for line in info:
        ws.append([line])
    ws.append([])
    fill, font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    header = []
    for c in columns:
        cell = WriteOnlyCell(ws, value=c.title)
        cell.font, cell.fill, cell.alignment = font, fill, Alignment(vertical="center")
        header.append(cell)
    ws.append(header)
    done = 0
    try:
        for chunk in chunks:
            for r in chunk:
                out = []
                for c in columns:
                    v = c.get(r)
                    if isinstance(v, datetime):
                        v = to_local(v).replace(tzinfo=None)
                    if c.kind in FORMATS and v is not None:
                        cell = WriteOnlyCell(ws, value=v)
                        cell.number_format = FORMATS[c.kind]
                        out.append(cell)
                    else:
                        out.append(v)
                ws.append(out)
            done += len(chunk)
            if progress is not None and progress(done, total) is False:
                raise _Cancelled
        last = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A{head}:{last}{head + max(done, 1)}"
        wb.save(path)
    except _Cancelled:
        Path(path).unlink(missing_ok=True)
        raise ExportCancelled() from None
    return done


class _Cancelled(Exception):
    pass


class ExportCancelled(Exception):
    """The person pressed Cancel during an export."""


def export_materials(svc: MaterialsService, path: Path, scope: str, f: MaterialFilter | None = None,
                     ids: list[int] | None = None, filters_text: str = "", company: str = "",
                     progress: Callable[[int, int | None], bool] | None = None) -> int:
    """Write the Materials list for one of the three scopes. Returns the number of rows."""
    f = (f or MaterialFilter()).copy()
    stamp = f"{company + ' - ' if company else ''}Exported {datetime.now():%d/%m/%Y %H:%M}"
    if scope == SELECTED:
        ids = list(ids or [])
        chunks = (svc.rows_for(ids[i:i + 1000], f.site_id) for i in range(0, len(ids), 1000))
        total, what = len(ids), f"Selected rows: {len(ids)}"
    elif scope == ALL:
        f = MaterialFilter(status="ALL")
        total = svc.count(f)
        chunks, what = svc.iter_rows(f), f"All materials (active and archived): {total}"
    else:
        total = svc.count(f)
        chunks, what = svc.iter_rows(f), f"Current view: {total} - {filters_text or 'no filters'}"
    columns = list(MATERIAL_COLUMNS)
    if f.site_id and scope != ALL:
        columns = PROJECT_COLUMNS + [c for c in MATERIAL_COLUMNS
                                     if c.title in ("Barcode", "Category", "On Hand", "Bins", "Stock Status")]
    return write_workbook(Path(path), "Materials", [stamp, what], columns, chunks, total, progress)


def export_project(svc: MaterialsService, site_id: int, path: Path, company: str = "",
                   progress: Callable[[int, int | None], bool] | None = None) -> int:
    """Projects -> Project -> Materials -> Export to Excel (spec section 22)."""
    summary = svc.project_summary(site_id)
    site = summary.site
    f = MaterialFilter(site_id=site_id, status="ALL")
    info = [
        f"Project: {site.project_number} - {site.name}",
        f"Customer: {site.customer or '-'}   Project manager: {site.project_manager or '-'}   "
        f"Status: {site.status.replace('_', ' ').title()}",
        f"{summary.materials} materials · Required {summary.required.normalize():f} · Reserved "
        f"{summary.reserved.normalize():f} · Issued {summary.issued.normalize():f} · Returned "
        f"{summary.returned.normalize():f} · Net issued {summary.net.normalize():f}",
        f"{company + ' - ' if company else ''}Exported {datetime.now():%d/%m/%Y %H:%M}",
    ]
    return write_workbook(Path(path), f"{site.project_number} Materials", info, PROJECT_COLUMNS,
                          svc.iter_rows(f), summary.materials, progress)


__all__ = ["ALL", "CURRENT", "MATERIAL_COLUMNS", "PROJECT_COLUMNS", "SELECTED", "ExportCancelled",
           "export_materials", "export_project", "write_workbook"]
