"""Phase 1, Stage 6 - Excel import of the existing materials list (with photos) and Excel export.

Every workbook here is built with openpyxl the way real ones are made: a title above the header,
pictures floating over cells, Excel 365 "Place in Cell" pictures, file paths, hyperlinks, bare file
names and web links."""

from __future__ import annotations

import io
import re
import zipfile
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from wms.core.exceptions import PermissionDeniedError, ValidationError
from wms.services.excel_import import (
    CREATE,
    DUPLICATE,
    ERROR,
    NEW,
    SKIP,
    UPDATE,
    ImportOptions,
    guess_mapping,
)

D = Decimal
HEAD = ["Code", "Name / Description", "Product Number", "Product Photo", "Supplier", "Unit", "Min"]


def png(colour=(200, 30, 30), size=(24, 24), fmt="PNG") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, colour).save(buf, fmt)
    return buf.getvalue()


def workbook(path, rows, title=True, images=None):
    """rows: lists under HEAD. images: {excel_row: png bytes} floating over column D."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    if title:
        ws.append(["Workshop materials list - October"])
        ws.append([])
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    for row, data in (images or {}).items():
        img = XLImage(io.BytesIO(data))
        ws.add_image(img, f"D{row}")
    wb.create_sheet("Notes").append(["nothing here"])
    wb.save(path)
    return path


def add_in_cell_picture(path, cell: str, data: bytes) -> None:
    """Turn `cell` of the first sheet into an Excel 365 "Place in Cell" picture (rich value)."""
    src = zipfile.ZipFile(path)
    parts = {n: src.read(n) for n in src.namelist()}
    src.close()
    sheet = "xl/worksheets/sheet1.xml"
    xml = parts[sheet].decode()
    xml, n = re.subn(rf'<c r="{cell}"[^>]*?(/>|>.*?</c>)', f'<c r="{cell}" t="e" vm="1"><v>#VALUE!</v></c>', xml)
    assert n == 1
    parts[sheet] = xml.encode()
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    parts["xl/metadata.xml"] = (
        f'<metadata {ns} xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
        '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000"/></metadataTypes>'
        '<futureMetadata name="XLRICHVALUE" count="1"><bk><extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}">'
        '<xlrd:rvb i="0"/></ext></extLst></bk></futureMetadata>'
        '<valueMetadata count="1"><bk><rc t="1" v="0"/></bk></valueMetadata></metadata>').encode()
    rd = 'xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"'
    parts["xl/richData/rdrichvaluestructure.xml"] = (
        f'<rvStructures {rd} count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/>'
        '<k n="CalcOrigin" t="i"/></s></rvStructures>').encode()
    parts["xl/richData/rdrichvalue.xml"] = (
        f'<rvData {rd} count="1"><rv s="0"><v>0</v><v>5</v></rv></rvData>').encode()
    parts["xl/richData/richValueRel.xml"] = (
        b'<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><rel r:id="rId1"/>'
        b'</richValueRels>')
    parts["xl/richData/_rels/richValueRel.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        b'Target="../media/incell1.png"/></Relationships>')
    parts["xl/media/incell1.png"] = data
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in parts.items():
            z.writestr(name, body)


def options(info, **kw) -> ImportOptions:
    sheet = info.sheet("Materials")
    return ImportOptions(sheet=sheet.name, header_row=sheet.header_row, mapping=dict(sheet.guess), **kw)


def photos_of(svc, product_id):
    return svc.attachments.list_for("PRODUCT", product_id)


# --- reading the workbook ----------------------------------------------------------------------

def test_header_detection_and_column_guess(svc, tmp_path):
    path = workbook(tmp_path / "m.xlsx", [["H3311", "Egger Halifax oak", "H3311 ST10", "", "Egger", "SHEET", 5]])
    info = svc.excel_import.inspect(path)
    assert [s.name for s in info.sheets] == ["Materials", "Notes"]
    sheet = info.sheet("Materials")
    assert sheet.header_row == 3  # under the title line and a blank line
    assert sheet.guess == {"internal_code": 0, "name": 1, "supplier_code": 2, "photo": 3, "supplier": 4,
                           "unit": 5, "min_stock": 6}
    assert sheet.samples[0][:2] == ["H3311", "Egger Halifax oak"]
    assert guess_mapping(["Item code", "Description", "Price"]) == {"internal_code": 0, "name": 1, "unit_cost": 2}


def test_photo_sources_are_counted(svc, tmp_path):
    folder = tmp_path / "pics"
    folder.mkdir()
    (folder / "u708.jpg").write_bytes(png(fmt="JPEG"))
    rows = [["A1", "Floating", "", "", "", "SHEET", ""],
            ["A2", "In cell", "", "x", "", "SHEET", ""],
            ["A3", "Path", "", str(folder / "u708.jpg"), "", "SHEET", ""],
            ["A4", "Name", "", "u708.jpg", "", "SHEET", ""],
            ["A5", "Web", "", "https://example.com/a.jpg", "", "SHEET", ""]]
    path = workbook(tmp_path / "m.xlsx", rows, images={4: png()})
    add_in_cell_picture(path, "D5", png((0, 0, 255)))
    sheet = svc.excel_import.inspect(path).sheet("Materials")
    assert (sheet.floating_images, sheet.cell_images, sheet.paths, sheet.filenames, sheet.web_links) == (1, 1, 1, 1, 1)
    assert "1 picture(s) placed in cells" in sheet.photo_summary


# --- the full import ---------------------------------------------------------------------------

def test_import_with_every_kind_of_photo(svc, sample, tmp_path):
    folder = tmp_path / "pics"
    (folder / "sub").mkdir(parents=True)
    (folder / "sub" / "W100.JPG").write_bytes(png((10, 200, 10), fmt="JPEG"))
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    (elsewhere / "u708.png").write_bytes(png((10, 10, 10)))
    linked = elsewhere / "linked.png"
    linked.write_bytes(png((50, 50, 50)))
    rows = [["H3311", "Egger Halifax oak 18mm", "H3311 ST10", "", "Egger UK Ltd", "SHEET", 5],
            ["U708", "Light grey 18mm", "U708 ST9", str(elsewhere / "u708.png"), "Egger UK Ltd", "SHEET", ""],
            ["W100", "Premium white", "W1000", "w100.jpg", "Kronospan", "sheet", "2"],
            ["H1145", "Bardolino oak", "H1145", "x", "Egger UK Ltd", "SHEET", ""],
            ["LINK-1", "Hyperlinked photo", "", "see photo", "", "M", ""],
            ["WEB-1", "Web photo", "", "https://example.com/p.jpg", "", "EA", ""]]
    path = workbook(tmp_path / "materials.xlsx", rows, images={4: png()})
    wb = load_workbook(path)
    wb["Materials"]["D8"].hyperlink = str(linked)
    wb.save(path)
    add_in_cell_picture(path, "D7", png((0, 0, 255)))
    info = svc.excel_import.inspect(path)
    plan = svc.excel_import.prepare(path, options(info, photo_folder=folder))
    by = {r.code: r for r in plan.rows}
    assert {r.action for r in plan.rows} == {CREATE}
    assert by["H3311"].photo.kind == "EMBEDDED"
    assert by["U708"].photo.kind == "PATH"
    assert by["W100"].photo.kind == "FILENAME" and by["W100"].photo.path.name == "W100.JPG"
    assert by["LINK-1"].photo.kind == "HYPERLINK"
    assert by["WEB-1"].photo is None and "web address" in by["WEB-1"].problems
    assert "new supplier 'Kronospan'" in by["W100"].problems
    result = svc.excel_import.run(plan)
    assert (result.created, result.updated, result.failed_rows, result.photo_problems) == (6, 0, 0, [])
    assert result.suppliers_created == ["Kronospan"]
    assert result.backup.exists() and result.backup.name.endswith("_safety.zip")
    h = svc.search.exact_product("H3311")
    assert h.supplier.name == "Egger UK Ltd" and h.supplier_code == "H3311 ST10" and h.min_stock == D(5)
    assert svc.search.exact_product("W100").supplier.name == "Kronospan"
    for code in ("H3311", "U708", "W100", "LINK-1"):
        files = photos_of(svc, svc.search.exact_product(code).id)
        assert len(files) == 1 and files[0].is_primary, code
    # Photos live in the WMS store: deleting the workbook and the picture folder loses nothing
    path.unlink()
    for f in (folder / "sub" / "W100.JPG", elsewhere / "u708.png", linked):
        f.unlink()
    for code in ("H3311", "U708", "W100", "LINK-1"):
        att = photos_of(svc, svc.search.exact_product(code).id)[0]
        assert svc.attachments.path_of(att).is_file()
    # Audit trail
    log = svc.audit.search(action="EXCEL_IMPORT")
    assert len(log) == 1 and "6 created" in log[0].summary


def test_in_cell_picture(svc, tmp_path):
    rows = [["H1145", "Bardolino oak", "", "x", "", "SHEET", ""]]
    path = workbook(tmp_path / "m.xlsx", rows)
    add_in_cell_picture(path, "D4", png((0, 0, 255)))
    plan = svc.excel_import.prepare(path, options(svc.excel_import.inspect(path)))
    assert plan.rows[0].photo.kind == "IN_CELL"
    svc.excel_import.run(plan)
    p = svc.search.exact_product("H1145")
    att = photos_of(svc, p.id)
    assert len(att) == 1
    assert svc.attachments.path_of(att[0]).read_bytes() == png((0, 0, 255))


def test_validation_and_duplicates(svc, sample, tmp_path):
    rows = [["MDF-18", "MDF 18mm (from Excel)", "", "", "", "SHEET", "12"],  # already in the WMS
            ["NEW-1", "New board", "", "", "", "SHEET", ""],
            ["new-1", "Same code again", "", "", "", "SHEET", ""],  # duplicate inside the file
            ["", "No code", "", "", "", "SHEET", ""],
            ["BAD-UNIT", "Bad unit", "", "", "", "PALLETS", ""],
            ["BAD-MIN", "Bad min", "", "", "", "SHEET", "lots"],
            ["HALF", "Half sheet", "", "", "", "SHEET", "1.5"],
            ["TWIN", "Same supplier code", "EG-MDF18", "", "Egger UK Ltd", "SHEET", ""]]
    path = workbook(tmp_path / "m.xlsx", rows)
    info = svc.excel_import.inspect(path)
    plan = svc.excel_import.prepare(path, options(info))
    by = {r.row: r for r in plan.rows}
    assert by[4].action == SKIP and by[4].existing_id == sample.mdf.id and by[4].ok
    assert by[5].action == CREATE
    assert by[4].status == DUPLICATE and by[5].status == NEW
    assert "also on row 5" in by[6].problems and by[6].status == DUPLICATE and by[6].action == SKIP
    assert plan.choices(by[6]) == [SKIP]  # a duplicate inside the file can never be imported twice
    assert "no code" in by[7].problems and by[7].status == ERROR
    assert "unknown unit 'PALLETS'" in by[8].problems
    assert "is not a number" in by[9].problems
    assert "must be a whole number" in by[10].problems
    # same product number as MDF-18 (Egger): found as a possible duplicate, NOT created silently
    assert by[11].ok and "possible duplicate" in by[11].problems and by[11].existing_id == sample.mdf.id
    assert by[11].match_kind == "PRODUCT_NO" and by[11].action == SKIP and by[11].status == DUPLICATE
    assert plan.choices(by[11]) == [SKIP, UPDATE, CREATE]
    assert ("Name", "MDF 18mm 2440x1220", "MDF 18mm (from Excel)") in by[4].changes
    c = plan.counts()
    assert (c["create"], c["update"], c["duplicate"], c["skip"], c["errors"]) == (1, 0, 3, 0, 4)
    by[11].action = CREATE  # the person decides it really is a different material
    # Update instead of skip: only mapped + filled columns change, stock and history stay
    sample.receive(sample.mdf, 7)
    plan = svc.excel_import.prepare(path, options(info, on_existing=UPDATE))
    assert next(r for r in plan.rows if r.row == 4).action == UPDATE
    result = svc.excel_import.run(plan)
    assert (result.created, result.updated) == (1, 1)
    mdf = svc.products.get(sample.mdf.id)
    assert (mdf.name, mdf.min_stock) == ("MDF 18mm (from Excel)", D(12))
    assert mdf.supplier_code == "EG-MDF18" and mdf.unit_cost == D("24.50")  # not mapped / empty: unchanged
    assert svc.products.row(mdf.id).on_hand == D(7)


def test_import_is_all_or_nothing(svc, sample, tmp_path, monkeypatch):
    rows = [["OK-1", "First", "", "", "", "SHEET", ""], ["OK-2", "Second", "", "", "", "SHEET", ""]]
    path = workbook(tmp_path / "m.xlsx", rows, images={4: png()})
    plan = svc.excel_import.prepare(path, options(svc.excel_import.inspect(path)))
    before = len(list(svc.attachments.store.rglob("*.*")))
    from wms.services import products as products_mod
    real = products_mod.ProductService.save_in

    def fail_second(self, s, product_id, data, expected_version=None):
        if data["internal_code"] == "OK-2":
            raise ValidationError("disk on fire")
        return real(self, s, product_id, data, expected_version)

    monkeypatch.setattr(products_mod.ProductService, "save_in", fail_second)
    with pytest.raises(ValidationError, match="Row 5 \\(OK-2\\): disk on fire - nothing was imported"):
        svc.excel_import.run(plan)
    assert svc.search.exact_product("OK-1") is None
    assert len(list(svc.attachments.store.rglob("*.*"))) == before  # staged photo removed again
    assert svc.audit.search(action="EXCEL_IMPORT") == []


def test_no_import_without_backup(svc, tmp_path, monkeypatch):
    path = workbook(tmp_path / "m.xlsx", [["X-1", "Thing", "", "", "", "EA", ""]])
    plan = svc.excel_import.prepare(path, options(svc.excel_import.inspect(path)))
    from wms.services.backup import BackupService

    def broken(self, note):
        raise OSError("backup drive full")

    monkeypatch.setattr(BackupService, "create_safety", broken)
    with pytest.raises(ValidationError, match="pre-import backup failed"):
        svc.excel_import.run(plan)
    assert svc.search.exact_product("X-1") is None


def test_mapping_rules(svc, tmp_path):
    path = workbook(tmp_path / "m.xlsx", [["X-1", "Thing", "", "", "", "", ""]])
    info = svc.excel_import.inspect(path)
    with pytest.raises(ValidationError, match="Choose the column for: Code"):
        svc.excel_import.prepare(path, ImportOptions("Materials", 3, {"name": 1}))
    with pytest.raises(ValidationError, match="unit of measure"):
        o = options(info)
        o.mapping.pop("unit")
        svc.excel_import.prepare(path, o)
    units = {u.code: u for u in svc.reference.units()}
    plan = svc.excel_import.prepare(path, options(info, default_unit_id=units["EA"].id))
    assert plan.rows[0].ok  # empty Unit cell -> default unit
    bad = tmp_path / "old.xls"
    bad.write_bytes(b"not really")
    with pytest.raises(ValidationError, match="Save As"):
        svc.excel_import.inspect(bad)


def test_operative_cannot_import(svc, tmp_path):
    path = workbook(tmp_path / "m.xlsx", [["X-1", "Thing", "", "", "", "EA", ""]])
    roles = {r.code: r for r in svc.users.list_roles()}
    svc.users.create_user({"username": "op", "full_name": "Op", "role_id": roles["OPERATIVE"].id,
                           "must_change_password": False}, "secret1")
    svc.auth.login("op", "secret1")
    with pytest.raises(PermissionDeniedError):
        svc.excel_import.inspect(path)


def test_damaged_photo_is_reported_in_the_preview(svc, tmp_path):
    (tmp_path / "fake.jpg").write_bytes(png())  # a PNG renamed to .jpg
    path = workbook(tmp_path / "m.xlsx", [["X-1", "Thing", "", "fake.jpg", "", "EA", ""]])
    plan = svc.excel_import.prepare(path, options(svc.excel_import.inspect(path)))
    row = plan.rows[0]
    assert row.ok and row.photo is None and "not a valid JPG" in row.problems


# --- Excel export -----------------------------------------------------------------------------

def test_product_export_formatting_and_round_trip(svc, sample, tmp_path):
    from wms.exporters.export import export_xlsx
    sample.receive(sample.mdf, 12)
    cat = svc.reference.save_category(None, "Boards")
    sub = svc.reference.save_category(None, "MDF", cat.id)
    p = svc.products.get(sample.mdf.id)
    svc.products.save(p.id, {"item_type": p.item_type, "internal_code": p.internal_code, "name": p.name,
                             "unit_id": p.unit_id, "supplier_id": p.supplier_id, "supplier_code": p.supplier_code,
                             "default_bin_id": p.default_bin_id, "min_stock": "10", "unit_cost": "24.50",
                             "category_id": sub.id}, p.version)
    everything = svc.reports.product_export()
    assert [r[0] for r in everything.rows] == ["CAB-600", "EDGE-W22", "MDF-18"]
    picked = svc.reports.product_export([sample.edge.id, sample.mdf.id], "Selected products")
    assert [r[0] for r in picked.rows] == ["EDGE-W22", "MDF-18"]
    assert picked.filters_text == "Selected products - 2 product(s)"
    path = export_xlsx(picked, tmp_path / "products.xlsx")
    ws = load_workbook(path)["Products"]
    header = [c.value for c in ws[5]]
    assert header[:6] == ["Code", "Name", "Type", "Description", "Product Number", "Supplier"]
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref.startswith("A5:")
    on_hand = header.index("On Hand") + 1
    mdf = next(r for r in ws.iter_rows(min_row=6) if r[0].value == "MDF-18")
    assert mdf[on_hand - 1].value == 12 and mdf[on_hand - 1].number_format == "#,##0.###"  # a number, not text
    assert mdf[header.index("Unit Cost")].number_format == "£#,##0.00"
    assert mdf[header.index("Created")].number_format == "dd/mm/yyyy hh:mm"
    assert ws.column_dimensions["B"].width > ws.column_dimensions["J"].width  # Name wider than Unit
    # The exported file goes straight back in: same products, nothing changes, no errors
    info = svc.excel_import.inspect(path)
    sheet = info.sheet("Products")
    assert sheet.header_row == 5
    assert {"internal_code", "name", "description", "supplier_code", "supplier", "barcode", "brand", "category",
            "unit", "min_stock", "reorder_level", "reorder_qty", "unit_cost", "default_bin", "notes"} <= set(sheet.guess)
    plan = svc.excel_import.prepare(path, ImportOptions("Products", 5, sheet.guess, on_existing=UPDATE))
    assert [(r.code, r.action, r.errors) for r in plan.rows] == [("EDGE-W22", UPDATE, []), ("MDF-18", UPDATE, [])]
    svc.excel_import.run(plan)
    after = svc.products.get(sample.mdf.id)
    assert (after.category_id, after.min_stock, after.unit_cost, after.default_bin_id) == (
        sub.id, D(10), D("24.50"), sample.bin1.id)
