"""Excel import of materials and project material lists: statuses, differences, supplier and project
resolution, mapping profiles, the template, the error report - and that stock is never touched."""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from wms.core.exceptions import ValidationError
from wms.services.excel_import import (
    CREATE,
    DUPLICATE,
    ERROR,
    IGNORE,
    MATCH,
    NEW,
    SKIP,
    SKIP_ROW,
    TEMPLATE_COLUMNS,
    UPDATE,
    ImportCancelled,
    ImportOptions,
    copy_template,
    export_error_report,
    supplier_key,
)

D = Decimal
HEAD = ["Code", "Material Name", "Product Number", "Supplier", "Unit", "Project", "Required Quantity"]


def book(path, rows, head=HEAD):
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(head)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def plan_for(svc, path, **kw):
    info = svc.excel_import.inspect(path)
    sheet = info.sheet("Materials")
    return svc.excel_import.prepare(path, ImportOptions(sheet.name, sheet.header_row, dict(sheet.guess), **kw))


def stock_snapshot(svc):
    return {r.product.internal_code: r.on_hand for r in svc.products.list(include_inactive=True)}


def test_supplier_names_are_normalised():
    assert supplier_key("EGGER") == supplier_key("Egger UK Ltd") == supplier_key("egger (uk) limited") == "EGGER"
    assert supplier_key("Kronospan Ltd") != supplier_key("Egger")


def test_suppliers_are_matched_or_listed_never_duplicated(svc, sample, tmp_path):
    path = book(tmp_path / "s.xlsx", [
        ["H3311", "Halifax Oak", "H3311 ST12", "EGGER", "SHEET", "", ""],  # = Egger UK Ltd
        ["K001", "Kronospan white", "K001 PE", "Kronospan", "SHEET", "", ""],
        ["K002", "Kronospan black", "K002 PE", "KRONOSPAN LTD", "SHEET", "", ""],  # same new supplier
        ["X1", "Unknown maker", "", "Nobody & Co", "SHEET", "", ""]])
    plan = plan_for(svc, path)
    egger = plan.suppliers[supplier_key("EGGER")]
    assert (egger.action, egger.match_id, egger.suggested) == (MATCH, sample.supplier.id, True)
    krono = plan.suppliers[supplier_key("Kronospan")]
    assert krono.action == CREATE and krono.rows == 2
    plan.suppliers[supplier_key("Nobody & Co")].action = IGNORE
    result = svc.excel_import.run(plan)
    assert result.suppliers_created == ["Kronospan"]  # one supplier for both spellings
    names = [s.name for s in svc.reference.suppliers(include_inactive=True)]
    assert names.count("Kronospan") == 1 and "Nobody & Co" not in names and "EGGER" not in names
    oak = svc.search.exact_product("H3311")
    assert oak.supplier_id == sample.supplier.id
    assert svc.search.exact_product("X1").supplier_id is None


def test_project_column_match_create_ignore_skip_and_closed(svc, sample, tmp_path):
    closed = svc.sites.save(None, {"project_number": "P-0900", "name": "Old Job", "status": "COMPLETED"})
    sample.receive(sample.mdf, 30)
    before = stock_snapshot(svc)
    path = book(tmp_path / "p.xlsx", [
        ["MDF-18", "MDF 18mm 2440x1220", "EG-MDF18", "Egger UK Ltd", "SHEET", "P-1001", "12"],  # existing material
        ["N-1", "New panel", "", "", "SHEET", "Riverside Apartments", "4"],  # project by name
        ["N-2", "Tribeca panel", "", "", "SHEET", "Tribeca", "8"],  # unknown project
        ["N-3", "Skipped panel", "", "", "SHEET", "Nowhere", "1"],
        ["N-4", "Closed job panel", "", "", "SHEET", "P-0900", "2"],
        ["N-5", "Bad qty", "", "", "SHEET", "P-1001", "lots"]])
    plan = plan_for(svc, path)
    p1001 = plan.projects["P-1001"]
    assert (p1001.action, p1001.match_id) == (MATCH, sample.site.id)
    assert plan.projects["RIVERSIDE APARTMENTS"].match_id == sample.site.id
    assert plan.projects["TRIBECA"].action == IGNORE  # never created without the person's choice
    assert plan.projects["P-0900"].closed and plan.projects["P-0900"].action == IGNORE
    by = {r.code: r for r in plan.rows}
    assert by["N-5"].status == ERROR and "not a number" in by["N-5"].problems
    assert by["MDF-18"].status == DUPLICATE and plan.links_project(by["MDF-18"])  # listed, master data untouched
    plan.projects["TRIBECA"].action = CREATE
    plan.projects["NOWHERE"].action = SKIP_ROW
    assert not plan.links_project(by["N-3"]) and by["N-3"] not in plan.to_import
    result = svc.excel_import.run(plan)
    assert result.projects_created == ["TRIBECA - Tribeca"]
    assert result.project_lines == 3 and result.created == 3  # N-1, N-2, N-4 (N-4 without its closed project)
    assert svc.search.exact_product("N-3") is None
    rows = {r.product.internal_code: r for r in svc.materials.rows_for(
        [svc.search.exact_product(c).id for c in ("MDF-18", "N-1")], sample.site.id)}
    assert rows["MDF-18"].required == D(12) and rows["N-1"].required == D(4)
    tribeca = next(s for s in svc.sites.list() if s.project_number == "TRIBECA")
    assert svc.materials.project_summary(tribeca.id).required == D(8)
    assert svc.materials.project_summary(closed.id).materials == 0
    # Import never changes stock: everything that existed has the same quantity, new ones have none
    after = stock_snapshot(svc)
    assert {k: v for k, v in after.items() if k in before} == before
    assert all(after[c] == 0 for c in ("N-1", "N-2", "N-4"))


def test_import_into_project_requires_quantity_and_open_project(svc, sample, tmp_path):
    path = book(tmp_path / "list.xlsx", [["MDF-18", "MDF 18mm 2440x1220", "", "", "SHEET", "", "20"],
                                         ["EDGE-W22", "ABS edging white 22mm", "", "", "M", "", "150.5"],
                                         ["CAB-600", "Base cabinet 600", "", "", "EA", "", ""]])
    info = svc.excel_import.inspect(path)
    sheet = info.sheet("Materials")
    mapping = {k: v for k, v in sheet.guess.items() if k != "required_qty"}
    with pytest.raises(ValidationError, match="Required Quantity"):
        svc.excel_import.prepare(path, ImportOptions("Materials", 1, mapping, project_id=sample.site.id))
    plan = plan_for(svc, path, project_id=sample.site.id)
    assert plan.project_label == "P-1001 - Riverside Apartments"
    assert plan.counts()["project_lines"] == 2  # CAB-600 has no quantity: a note, not on the list
    assert "no Required Quantity" in next(r for r in plan.rows if r.code == "CAB-600").problems
    result = svc.excel_import.run(plan)
    assert (result.created, result.updated, result.project_lines) == (0, 0, 2)
    summ = svc.materials.project_summary(sample.site.id)
    assert (summ.materials, summ.required) == (2, D("170.5"))
    assert svc.products.movements(sample.edge.id) == []  # no Goods Out, no movement of any kind
    svc.sites.set_status(sample.site.id, "COMPLETED")
    with pytest.raises(ValidationError, match="can no longer be changed"):
        plan_for(svc, path, project_id=sample.site.id)


def test_update_shows_differences_and_apply_to_similar(svc, sample, tmp_path):
    path = book(tmp_path / "u.xlsx", [["MDF-18", "MDF 18mm FSC", "EG-MDF18", "Egger UK Ltd", "SHEET", "", ""],
                                      ["EDGE-W22", "ABS edging white 22mm", "", "", "M", "", ""],
                                      ["MDF19", "Same board other code", "EG-MDF18", "", "SHEET", "", ""]])
    plan = plan_for(svc, path)
    by = {r.code: r for r in plan.rows}
    assert by["MDF-18"].changes == [("Name", "MDF 18mm 2440x1220", "MDF 18mm FSC")]
    assert by["EDGE-W22"].changes == [] and by["EDGE-W22"].changes_text == "no changes"
    assert by["MDF19"].match_kind == "PRODUCT_NO" and by["MDF19"].existing_id == sample.mdf.id
    assert {by[c].status for c in by} == {DUPLICATE}
    by["MDF-18"].action = UPDATE
    assert plan.apply_to_similar(by["MDF-18"], UPDATE) == 2  # both code matches; not the product-number one
    assert by["EDGE-W22"].action == UPDATE and by["MDF19"].action == SKIP
    # choices survive a re-check of the preview
    again = svc.excel_import.prepare(plan.source, plan.options, plan)
    assert [r.action for r in again.rows] == [UPDATE, UPDATE, SKIP]
    svc.excel_import.run(again)
    assert svc.products.get(sample.mdf.id).name == "MDF 18mm FSC"
    assert svc.search.exact_product("MDF19") is None


def test_cancel_rolls_back_everything(svc, sample, tmp_path):
    rows = [[f"C-{i:03}", f"Board {i}", "", "", "SHEET", "", ""] for i in range(60)]
    plan = plan_for(svc, book(tmp_path / "c.xlsx", rows))
    seen = []

    def progress(done, total):
        seen.append(done)
        return done < 30  # press Cancel half way

    with pytest.raises(ImportCancelled):
        svc.excel_import.run(plan, progress)
    assert seen and all(svc.search.exact_product(f"C-{i:03}") is None for i in range(60))
    done = []
    result = svc.excel_import.run(plan, lambda d, t: done.append((d, t)) or True)
    assert result.created == 60 and done[-1] == (60, 60)


def test_mapping_profiles_are_remembered_by_column_title(svc, sample, tmp_path):
    head = ["Our Ref", "Decor", "Maker", "UOM"]
    path = book(tmp_path / "odd.xlsx", [["Z1", "Odd board", "Egger UK Ltd", "SHEET"]], head=head)
    info = svc.excel_import.inspect(path)
    assert info.sheet("Materials").profile is None
    svc.excel_import.save_profile("Supplier list", head, {"internal_code": 0, "name": 1, "supplier": 2, "unit": 3})
    # next file: same titles, other order -> recognised automatically
    path2 = book(tmp_path / "odd2.xlsx", [["Odd board", "SHEET", "Z2", "Egger UK Ltd"]],
                 head=["Decor", "UOM", "Our Ref", "Maker"])
    sheet = svc.excel_import.inspect(path2).sheet("Materials")
    assert sheet.profile == "Supplier list"
    assert sheet.guess == {"internal_code": 2, "name": 0, "supplier": 3, "unit": 1}
    assert [p[1] for p in svc.excel_import.profiles()] == ["Supplier list"]
    assert svc.search.saved_searches() == []  # profiles are not shown as saved searches
    pid = svc.excel_import.profiles()[0][0]
    svc.excel_import.delete_profile(pid)
    assert svc.excel_import.profiles() == []


def test_template_and_error_report(svc, sample, tmp_path):
    t = copy_template(tmp_path / "template.xlsx")
    wb = load_workbook(t)
    ws = wb["Materials"]
    titles = [c.value for c in ws[1]]
    assert titles[:9] == ["Code", "Material Name", "Product Number", "Supplier", "Photo", "Category", "Brand",
                          "Project", "Required Quantity"]
    assert ws.freeze_panes == "A2" and "How to fill in" in wb.sheetnames
    # the template's own titles are recognised without any manual mapping
    ws.append(["T-1", "Template board", "", "", "", "", "", "", ""])
    wb.save(t)
    sheet = svc.excel_import.inspect(t).sheet("Materials")
    assert {f for f, _ in TEMPLATE_COLUMNS} <= set(sheet.guess)
    path = book(tmp_path / "bad.xlsx", [["", "No code", "", "", "SHEET", "", ""],
                                        ["OK-1", "Fine", "", "", "SHEET", "", ""],
                                        ["OK-1", "Twice", "", "", "SHEET", "", ""]])
    plan = plan_for(svc, path)
    assert [r.status for r in plan.rows] == [ERROR, NEW, DUPLICATE]
    out = export_error_report(plan, tmp_path / "problems.xlsx")
    rep = load_workbook(out).active
    lines = [[c.value for c in r] for r in rep.iter_rows(min_row=5) if r[0].value]
    assert [(r[0], r[1]) for r in lines] == [(2, "ERROR"), (4, "DUPLICATE")]
    assert "no code" in lines[0][6] and "also on row 3" in lines[1][6]
