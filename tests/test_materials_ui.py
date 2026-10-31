"""The Materials screen and a project's Materials tab, driven like a person would."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt
from test_ui_smoke import _errors, ui  # noqa: F401 - the shared screen fixture

from wms.services.materials import MaterialFilter
from wms.ui.widgets import materials_view as mv

D = Decimal


def tick(view, *rows):
    for r in rows:
        view.model.setData(view.model.index(r, 0), Qt.CheckState.Checked.value, Qt.ItemDataRole.CheckStateRole)


def codes(view):
    return [r.product.internal_code for r in view.model.rows]


def page_of(win):
    win.go("products")
    page = win.page("products")
    page.on_show()
    return page


def test_filters_chips_count_and_clear(ui, sample):  # noqa: F811
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 25)
    page = page_of(win)
    assert page.title == "Materials" and win.nav.currentItem().text() == "Materials"
    v = page.view
    assert v.count.text() == "Showing 3 materials"  # boards, edging and the finished cabinet
    assert codes(v) == ["CAB-600", "EDGE-W22", "MDF-18"]  # by code
    v.supplier.set_value(sample.supplier.id)
    assert codes(v) == ["MDF-18"] and v.count.text() == "Showing 1 material"
    assert "supplier_id" in v.chip_buttons
    v.stock.setCurrentIndex(v.stock.findData("OUT_OF_STOCK"))
    assert codes(v) == [] and v.count.text() == "Showing 0 materials"
    v.chip_buttons["stock"].click()  # × on the chip
    assert codes(v) == ["MDF-18"] and v.stock.currentData() is None
    v.bin.set_value("NONE")
    assert codes(v) == []
    v.clear_all()
    assert v.filter.is_default() and v.supplier.value() is None and v.bin.value() is None
    assert not _errors(msgs), msgs


def test_more_filters_views_and_columns(ui, sample, monkeypatch):  # noqa: F811
    app, win, msgs, _ = ui
    page = page_of(win)
    v = page.view
    dlg = mv.MoreFiltersDialog(app, v.filter, v)
    dlg.has_supplier.setCurrentIndex(dlg.has_supplier.findData("WITHOUT"))
    dlg.apply()
    v.filter = dlg.result_filter
    v.refresh()
    assert "EDGE-W22" in codes(v) and "MDF-18" not in codes(v)
    assert v.more.text() == "More filters (1)" and "has_supplier" in v.chip_buttons
    assert v.save_view("No supplier")
    assert any(x.name == "No supplier" for x in app.services.materials.views())
    v.clear_all()
    view = next(x for x in app.services.materials.views() if x.name == "No supplier")
    v.set_filter(view.filter)
    assert v.filter.has_supplier == "WITHOUT"
    v._fill_views()  # the Views menu, as when the person opens it
    next(a for a in v.views_menu.actions() if a.text() == "Low Stock").trigger()
    assert v.filter.stock == "LOW_STOCK" and v.stock.currentData() == "LOW_STOCK"
    # columns: hide Supplier, it stays hidden next time the screen opens
    assert any(c.key == "supplier" for c in v.model.cols)
    v.toggle_column("supplier", False)
    assert not any(c.key == "supplier" for c in v.model.cols)
    again = mv.MaterialsView(app)
    assert not any(c.key == "supplier" for c in again.shown_columns())
    again.reset_columns()
    assert any(c.key == "supplier" for c in again.shown_columns())
    # project columns only appear with a project filter
    assert not any(c.project for c in v.model.cols)
    v.project.set_value(sample.site.id)
    assert any(c.key == "required" for c in v.model.cols)
    assert v.count.text().endswith("in Riverside Apartments")
    assert not _errors(msgs), msgs


def test_rows_load_as_you_scroll_and_sort(ui, sample):  # noqa: F811
    app, win, msgs, _ = ui
    units = {u.code: u for u in app.services.reference.units()}
    for i in range(450):
        app.services.products.save(None, {"item_type": "RAW_MATERIAL", "internal_code": f"BULK-{i:04}",
                                          "name": f"Bulk board {i}", "unit_id": units["SHEET"].id})
    page = page_of(win)
    v = page.view
    assert v.model.total >= 452 and len(v.model.rows) == mv.PAGE_SIZE  # first page only
    assert v.count.text() == f"Showing {v.model.total:,} materials"
    assert v.model.canFetchMore()
    v.model.fetchMore()
    assert len(v.model.rows) == 2 * mv.PAGE_SIZE
    v.load_all()
    assert len(v.model.rows) == v.model.total  # no artificial limit
    name_col = v.model.col_of("name")
    v._header_clicked(name_col)  # sort by Material A-Z
    assert v.filter.sort == "name" and not v.filter.descending
    names = [r.product.name for r in v.model.rows]
    assert names == sorted(names, key=str.lower)
    v._header_clicked(name_col)
    assert v.filter.descending
    # "Select all" ticks every matching material, not just the loaded ones
    v.search.setText("bulk")
    v.search.search.emit("bulk")
    v.select_all()
    assert len(v.checked_ids()) == 450 and v.sel_label.text() == "450 selected"
    assert not _errors(msgs), msgs


def test_bulk_change_supplier_and_export_current_view(ui, sample, tmp_path, monkeypatch):  # noqa: F811
    app, win, msgs, _ = ui
    page = page_of(win)
    v = page.view
    v.search.setText("edging")
    v.search.search.emit("edging")
    tick(v, 0)
    assert v.sel_bar.isVisibleTo(page)
    n = app.services.materials.bulk_set(v.checked_ids(), "supplier_id", sample.supplier.id)
    assert n == 1 and app.services.products.get(sample.edge.id).supplier_id == sample.supplier.id
    out = page.export_excel("CURRENT", str(tmp_path / "current.xlsx"))
    ws = load_workbook(out).active
    assert ws["A3"].value == "Current view: 1 - Search: edging"
    assert [r[0].value for r in ws.iter_rows(min_row=6)] == ["EDGE-W22"]
    assert ws.freeze_panes == "A6" and ws.auto_filter.ref.startswith("A5:")
    assert not _errors(msgs), msgs


def test_project_tab_add_set_remove_and_export(ui, sample, tmp_path, monkeypatch):  # noqa: F811
    app, win, msgs, opened = ui
    from wms.ui.dialogs.site_dialog import SiteDialog
    dlg = SiteDialog(app, sample.site.id, tab="materials")
    assert dlg.tabs.tabText(0) == "Overview" and dlg.tabs.tabText(1) == "Materials"
    assert dlg.tabs.currentWidget() is dlg.summary_tab
    v = dlg.materials
    assert v.count.text() == "0 materials on this project" and not v.project.isVisible()
    assert {c.key for c in v.model.cols} >= {"required", "proj_reserved", "issued", "returned", "net"}
    # put two materials on the list (as "Add materials..." does after the picker)
    rows = app.services.materials.rows_for([sample.mdf.id, sample.edge.id])
    pl = mv.ProjectListDialog(app, rows, sample.site.id, dlg)
    pl.set_quantity(0, "12")
    pl.set_quantity(1, "40.5")
    pl.save()
    assert pl.saved
    v.refresh()
    assert v.count.text() == "2 materials on this project"
    req = {r.product.internal_code: r.required for r in v.model.rows}
    assert req == {"MDF-18": D(12), "EDGE-W22": D("40.5")}
    dlg._fill_cards()
    assert dlg.cards["required"].value_label.text() == "52.5"
    assert app.services.products.movements(sample.mdf.id) == []  # the list never moves stock
    # export from the project's Materials tab
    out = v.export_project(str(tmp_path / "project.xlsx"))
    ws = load_workbook(out).active
    assert ws["A2"].value == "Project: P-1001 - Riverside Apartments"
    header = [c.value for c in ws[7]]
    assert header[:6] == ["Internal Code", "Material", "Product Number", "Supplier", "Unit", "Required"]
    # take one off the list (the confirmation is answered Yes by the fixture)
    v.load_all()
    edge_row = codes(v).index("EDGE-W22")
    tick(v, edge_row)
    assert v.remove_from_project()
    assert codes(v) == ["MDF-18"]
    assert not _errors(msgs), msgs


def test_import_project_list_wizard(ui, sample, tmp_path):  # noqa: F811
    app, win, msgs, _ = ui
    from wms.ui.dialogs.excel_import import ExcelImportDialog
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(["Code", "Material Name", "Supplier", "Unit", "Required Quantity"])
    ws.append(["MDF-18", "MDF 18mm 2440x1220", "EGGER", "SHEET", 30])
    ws.append(["NEW-9", "New worktop", "Worktops Direct", "EA", 2])
    path = tmp_path / "list.xlsx"
    wb.save(path)
    dlg = ExcelImportDialog(app, win, project_id=sample.site.id)
    assert dlg.project.value() == sample.site.id
    dlg.path.setText(str(path))
    dlg.next()
    dlg.next()
    assert dlg.map_combos["required_qty"].currentData() == 4
    assert dlg.save_profile("Project lists")
    dlg.next()
    assert dlg.stack.currentIndex() == 3
    assert "2 project material line(s)" in dlg.preview_summary.text()
    dlg.next()  # Resolve: EGGER is matched to Egger UK Ltd, Worktops Direct is new
    assert dlg.stack.currentIndex() == 4 and dlg.sup_table.rowCount() == 2
    dlg.set_supplier("Worktops Direct", "IGNORE")
    dlg.next()
    assert dlg.stack.currentIndex() == 5, msgs
    assert "Project material lines: 2" in dlg.done_text.text()
    summ = app.services.materials.project_summary(sample.site.id)
    assert (summ.materials, summ.required) == (2, D(32))
    assert app.services.search.exact_product("NEW-9").supplier_id is None
    assert app.services.products.get(sample.mdf.id).supplier_id == sample.supplier.id
    assert not _errors(msgs), msgs


def test_dashboard_style_preset(ui, sample):  # noqa: F811
    app, win, msgs, _ = ui
    page = page_of(win)
    page.apply_preset({"stock": "OUT_OF_STOCK"})
    assert page.view.stock.currentData() == "OUT_OF_STOCK"
    assert page.view.filter == MaterialFilter(stock="OUT_OF_STOCK")
    assert page.view.count.text() == "Showing 3 materials"  # nothing received yet
    assert not _errors(msgs), msgs


def test_sidebar_is_grouped_and_less_used_groups_fold(ui, sample):  # noqa: F811
    app, win, msgs, opened = ui
    groups = list(win.sections)
    assert groups == ["HOME", "WAREHOUSE", "OPERATIONS", "PROJECTS", "PURCHASING", "MORE", "ADMIN"]
    texts = {s: [i.text() for i in items] for s, items in win.section_items.items()}
    assert texts == {
        "HOME": ["Dashboard"],
        "WAREHOUSE": ["Materials", "Locations & Bins"],
        "OPERATIONS": ["Goods In", "Goods Out", "Transfer"],
        "PROJECTS": ["Projects"],
        "PURCHASING": ["Purchase Orders", "Deliveries"],
        "MORE": ["Stocktake", "Returns", "Finished Products", "Reports", "Import / Export", "Archived Records"],
        "ADMIN": ["Users", "Audit Log", "Backup & Restore", "Settings"],
    }
    # less frequent groups start folded; a click on the heading opens / closes them
    assert win.is_folded("MORE") and win.is_folded("ADMIN") and not win.is_folded("OPERATIONS")
    win._nav_clicked(win.sections["MORE"])
    assert not win.is_folded("MORE") and win.sections["MORE"].text() == "▾  MORE"
    win._nav_clicked(win.sections["MORE"])
    assert win.is_folded("MORE")
    # daily groups never fold
    win._nav_clicked(win.sections["OPERATIONS"])
    assert not win.is_folded("OPERATIONS") and win.sections["OPERATIONS"].text() == "OPERATIONS"
    # the folded state is remembered for the next start
    win._set_folded("ADMIN", False)
    assert win._folded_groups() == {"MORE"}
    win._set_folded("ADMIN", True)
    # going to a screen in a folded group opens that group
    win.go("users")
    assert not win.is_folded("ADMIN") and win.stack.currentWidget() is win.page("users")
    # entries that are actions run them and keep the current screen selected
    transfer = next(i for i in win.section_items["OPERATIONS"] if i.text() == "Transfer")
    win.nav.setCurrentItem(transfer)
    assert "TransferDialog" in opened and win.nav.currentItem().data(Qt.ItemDataRole.UserRole) == "users"
    backup = next(i for i in win.section_items["ADMIN"] if i.text() == "Backup & Restore")
    win.nav.setCurrentItem(backup)
    assert win.page("settings").tabs.tabText(win.page("settings").tabs.currentIndex()) == "Backup and restore"
    win._set_folded("MORE", False)
    archive = next(i for i in win.section_items["MORE"] if i.text() == "Archived Records")
    win.nav.setCurrentItem(archive)
    assert win.page("products").view.filter.status == "ARCHIVED"
    assert not _errors(msgs), msgs


def test_screens_without_a_sidebar_entry_open_in_context(ui, sample):  # noqa: F811
    """Spec 3: Reservations, Low Stock and Delivery Problems left the sidebar but are one click from
    where they belong, and the sidebar shows that place."""
    app, win, msgs, opened = ui
    keys = {i.data(Qt.ItemDataRole.UserRole) for items in win.section_items.values() for i in items}
    assert not keys & {"reservations", "low_stock", "discrepancies", "search", "movements"}
    for key, parent in (("reservations", "sites"), ("low_stock", "products"),
                        ("discrepancies", "delivery_notes"), ("movements", "reports")):
        win.go(parent)
        page = win.page(parent)
        link = next(page.actions.itemAt(i).widget() for i in range(page.actions.count())
                    if page.actions.itemAt(i).widget() is not None
                    and page.actions.itemAt(i).widget().text().startswith(win.titles[key].split()[0]))
        link.click()
        assert win.stack.currentWidget() is win.page(key), key
        assert win.nav.currentItem().data(Qt.ItemDataRole.UserRole) == parent
        # one click back
        back = win.page(key).actions.itemAt(0).widget()
        assert back.text().startswith("‹  Back to")
        back.click()
        assert win.stack.currentWidget() is page
    # Materials offers reordering while it shows low stock
    win.go("products")
    mats = win.page("products")
    assert mats.reorder_btn.isHidden()
    mats.apply_preset({"stock": "LOW_STOCK"})
    assert not mats.reorder_btn.isHidden()
    assert not _errors(msgs), msgs


def test_top_bar_quick_actions_and_user_menu(ui, sample):  # noqa: F811
    app, win, msgs, opened = ui
    assert [b.text() for b in win.quick.values()] == ["+  Goods In", "−  Goods Out", "⇄  Transfer"]
    assert "(F2)" in win.quick["goods_in"].toolTip()
    win.quick["goods_out"].click()
    assert win.stack.currentWidget() is win.page("goods_out")
    win.quick["transfer"].click()
    assert "TransferDialog" in opened
    actions = [a.text() for a in win.user_menu.actions() if a.text()]
    assert actions[-5:] == ["My Account", "Change Password...", "Settings", "Keyboard Shortcuts", "Log Out"]
    # no permanent account buttons in the top bar any more
    from PySide6.QtWidgets import QPushButton
    texts = [b.text() for b in win.findChildren(QPushButton)]
    assert "Change password" not in texts and "Log out" not in texts
    next(a for a in win.user_menu.actions() if a.text() == "Settings").trigger()
    assert win.stack.currentWidget() is win.page("settings")
    assert not _errors(msgs), msgs
