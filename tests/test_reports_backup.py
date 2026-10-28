"""Reports, exports, attachments and backup/restore."""

from __future__ import annotations

import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from wms.core.exceptions import ValidationError
from wms.db.engine import Database
from wms.exporters.export import export
from wms.services.goods_out import IssueLine
from wms.services.reports import ReportFilters
from wms.services.stock_ops import ReturnLineInput


@pytest.fixture
def busy(svc, sample):
    sample.receive(sample.mdf, 40, cost="23.90")
    sample.receive(sample.edge, 120, number="DN-2", cost="0.2")
    svc.goods_out.post({"issue_date": date.today(), "site_id": sample.site.id, "reference": "JOB-1"},
                       [IssueLine(sample.mdf.id, "12"), IssueLine(sample.edge.id, "30")])
    svc.stock.post_return({"return_type": "FROM_SITE", "return_date": date.today(), "site_id": sample.site.id},
                          [ReturnLineInput(sample.mdf.id, "2", sample.bin1.id)])
    svc.stock.adjust(sample.edge.id, sample.bin2.id, "89", "STOCK_COUNT", "count")
    svc.stock.transfer(sample.mdf.id, sample.bin1.id, sample.bin2.id, "5")
    return sample


def test_every_report_runs_and_exports(svc, busy, tmp_path):
    for rdef in svc.reports.definitions():
        result = svc.reports.run(rdef.key, ReportFilters())
        assert result.columns, rdef.key
        for ext in (".csv", ".xlsx", ".pdf"):
            out = export(result, tmp_path / f"{rdef.key}{ext}", "Test Ltd")
            assert out.stat().st_size > 0


def test_report_figures(svc, busy):
    stock = svc.reports.run("current_stock", ReportFilters())
    mdf = next(r for r in stock.rows if r[0] == "MDF-18")
    assert mdf[5] == Decimal("30")  # 40 - 12 + 2
    assert mdf[9] == Decimal("717.00")  # 30 x £23.90
    by_site = svc.reports.run("stock_by_site", ReportFilters(site_id=busy.site.id))
    mdf_site = next(r for r in by_site.rows if r[1] == "MDF-18")
    assert mdf_site[3:6] == [Decimal("12"), Decimal("2"), Decimal("10")]
    by_bin = svc.reports.run("stock_by_bin", ReportFilters(product_id=busy.mdf.id))
    assert {r[0].split(" / ")[-1]: r[4] for r in by_bin.rows} == {"B04": Decimal("25"), "B05": Decimal("5")}
    adj = svc.reports.run("adjustments", ReportFilters())
    assert adj.rows[0][5:8] == [Decimal("90"), Decimal("89"), Decimal("-1")]
    sup = svc.reports.run("supplier_deliveries", ReportFilters())
    assert sup.rows[0][1] == 2


def test_xlsx_numbers_are_numbers(svc, busy, tmp_path):
    out = export(svc.reports.run("current_stock", ReportFilters()), tmp_path / "s.xlsx")
    ws = load_workbook(out).active
    assert ws.cell(row=5, column=6).value == "On Hand"
    assert isinstance(ws.cell(row=6, column=6).value, (int, float))


def _png(path):
    # 1x1 transparent PNG
    path.write_bytes(bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
        "1f15c4890000000d49444154789c6360000002000154a24f5d00000000"
        "49454e44ae426082"))
    return path


def test_attachments_photo_and_document(svc, sample, tmp_path):
    photo = _png(tmp_path / "mdf.png")
    att = svc.attachments.add("PRODUCT", sample.mdf.id, photo)
    assert att.is_primary and att.kind == "PHOTO"
    assert svc.attachments.path_of(att).is_file()
    photo.unlink()  # original deleted from the Desktop - our copy survives
    assert svc.attachments.path_of(att).is_file()
    second = _png(tmp_path / "mdf2.png")
    assert svc.attachments.find_duplicates(second)  # same bytes as the first photo
    new = svc.attachments.replace_photo(att.id, second)
    photos = svc.attachments.list_for("PRODUCT", sample.mdf.id)
    assert [p.id for p in photos] == [new.id] and photos[0].is_primary
    bad = tmp_path / "virus.exe"
    bad.write_bytes(b"MZ")
    with pytest.raises(ValidationError):
        svc.attachments.add("PRODUCT", sample.mdf.id, bad)


def test_delivery_note_with_scan(svc, sample, tmp_path):
    from wms.services.goods_in import ReceiptLine
    pdf = tmp_path / "dn.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%fake\n")
    dn = svc.goods_in.post({"dn_number": "S1", "supplier_id": sample.supplier.id, "delivery_date": date.today()},
                           [ReceiptLine(sample.mdf.id, "1", sample.bin1.id)], attachments=[pdf])
    docs = svc.attachments.list_for("DELIVERY_NOTE", dn.id)
    assert len(docs) == 1 and docs[0].is_pdf


def test_backup_and_restore_roundtrip(svc, busy, cfg, tmp_path):
    photo = _png(tmp_path / "p.png")
    svc.attachments.add("PRODUCT", busy.mdf.id, photo)
    backup = svc.backup.create(note="before test")
    with zipfile.ZipFile(backup) as z:
        names = z.namelist()
    assert "wms.db" in names and "manifest.json" in names and any(n.startswith("attachments/") for n in names)
    assert svc.backup.verify(backup)["note"] == "before test"

    # Change data after the backup, then restore
    busy.receive(busy.mdf, 100, number="AFTER")
    safety = svc.backup.restore(backup)
    assert safety.exists()

    db2 = Database(cfg.resolved_database_url)
    db2.migrate()
    svc.ctx.db = db2
    svc.auth.login("admin", "admin")
    from wms.services.products import ProductService
    rows = {r.product.internal_code: r.on_hand for r in ProductService(svc.ctx).list()}
    assert rows["MDF-18"] == Decimal("30")
    db2.dispose()


def test_verify_rejects_non_backup(svc, tmp_path):
    f = tmp_path / "x.zip"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(ValidationError):
        svc.backup.verify(f)
