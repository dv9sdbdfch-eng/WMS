"""Materials workspace: one filter definition for the list, the count, the export and the project tab."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from wms.core.exceptions import ValidationError
from wms.exporters.materials import ALL, CURRENT, SELECTED, export_materials, export_project
from wms.services.goods_out import IssueLine
from wms.services.materials import MaterialFilter as F
from wms.services.stock_ops import ReturnLineInput

D = Decimal
TODAY = date.today()


def codes(svc, f: F) -> list[str]:
    return [r.product.internal_code for r in svc.materials.page(f.copy(), 0, 10_000)]


def archive(svc, ws, product):
    """Give the product history (so removing it archives instead of deleting), then remove it."""
    ws.receive(product, 1, ws.bin1, number=f"DN-A{product.id}")
    issue(svc, ws.site, product, 1, ws.bin1)
    assert not svc.products.remove(product.id, "no longer used").permanent


def issue(svc, site, product, qty, bin_=None):
    return svc.goods_out.post({"site_id": site.id, "issue_date": TODAY},
                              [IssueLine(product.id, str(qty), bin_.id if bin_ else None)])


@pytest.fixture
def ws(svc, sample):
    """Sample warehouse plus an EGGER oak board in two bins and a second supplier."""
    units = {u.code: u for u in svc.reference.units()}
    sample.egger = svc.reference.save_supplier(None, {"code": "EGGER", "name": "EGGER"})
    sample.oak = svc.products.save(None, {"item_type": "RAW_MATERIAL", "internal_code": "H3311",
                                          "name": "Halifax Oak Natural", "description": "18mm decor board, oak",
                                          "supplier_code": "H3311 ST12", "barcode": "5901234123457",
                                          "unit_id": units["SHEET"].id, "supplier_id": sample.egger.id,
                                          "default_bin_id": sample.bin1.id, "reorder_level": "40"})
    sample.grey = svc.products.save(None, {"item_type": "RAW_MATERIAL", "internal_code": "U708",
                                           "name": "Light Grey", "supplier_code": "U708 ST9",
                                           "unit_id": units["SHEET"].id, "supplier_id": sample.egger.id})
    sample.receive(sample.oak, 85, sample.bin1, number="DN-1047")
    sample.receive(sample.oak, 42, sample.bin2, number="DN-1048")
    sample.receive(sample.mdf, 20, number="DN-1049")
    return sample


# --- search ---------------------------------------------------------------------------------------

def test_search_is_partial_case_insensitive_and_trimmed(svc, ws):
    assert codes(svc, F(text="3311")) == ["H3311"]
    assert codes(svc, F(text="  h3311  ")) == ["H3311"]
    assert codes(svc, F(text="h 3311")) == ["H3311"]  # other spacing still finds the code
    assert codes(svc, F(text="OAK")) == ["H3311"]  # name / description
    assert codes(svc, F(text="st12")) == ["H3311"]  # product number
    assert codes(svc, F(text="123457")) == ["H3311"]  # barcode
    assert set(codes(svc, F(text="egger"))) == {"H3311", "U708", "MDF-18"}  # supplier name (Egger UK Ltd)
    assert codes(svc, F(text="grey egger")) == ["U708"]  # every word must match
    assert codes(svc, F(text="100%_x")) == []  # wildcards typed by the user are literal


# --- project association and the four project figures ---------------------------------------------

def test_project_list_reservation_issue_and_return_are_separate(svc, ws):
    site = ws.site
    before = svc.products.row(ws.oak.id).on_hand
    svc.materials.set_required(site.id, [(ws.oak.id, "50"), (ws.grey.id, "25")])
    # Adding to the project list changes no stock and creates no movement.
    assert svc.products.row(ws.oak.id).on_hand == before
    assert svc.products.movements(ws.grey.id) == []
    svc.reservations.reserve({"product_id": ws.oak.id, "site_id": site.id, "quantity": "30"})
    issue(svc, site, ws.oak, 20, ws.bin1)
    issue(svc, site, ws.mdf, 6)  # not on the list: still belongs to the project
    svc.stock.post_return({"return_type": "FROM_SITE", "return_date": TODAY, "site_id": site.id},
                          [ReturnLineInput(ws.mdf.id, "2", ws.bin1.id)])
    rows = {r.product.internal_code: r for r in svc.materials.page(F(site_id=site.id))}
    assert set(rows) == {"H3311", "U708", "MDF-18"}
    oak, grey, mdf = rows["H3311"], rows["U708"], rows["MDF-18"]
    assert (oak.required, oak.proj_reserved, oak.issued, oak.returned, oak.net) == (D(50), D(10), D(20), 0, D(20))
    assert (grey.required, grey.issued, grey.on_list) == (D(25), 0, True)
    assert (mdf.required, mdf.issued, mdf.returned, mdf.net, mdf.on_list) == (None, D(6), D(2), D(4), False)
    summary = svc.materials.project_summary(site.id)
    assert (summary.materials, summary.required, summary.issued, summary.returned) == (3, D(75), D(26), D(2))
    # replace=False adds to what is required
    svc.materials.set_required(site.id, [(ws.grey.id, "5")], replace=False)
    assert svc.materials.page(F(site_id=site.id, text="U708"))[0].required == D(30)
    svc.materials.remove_from_project(site.id, [ws.grey.id])
    assert "U708" not in codes(svc, F(site_id=site.id))


def test_reversed_goods_out_nets_off_in_issued(svc, ws):
    gi = issue(svc, ws.site, ws.oak, 12, ws.bin1)
    svc.goods_out.reverse(gi.id, "wrong project")
    row = svc.materials.page(F(site_id=ws.site.id))[0]
    assert (row.issued, row.net) == (0, 0)


def test_project_list_of_closed_project_cannot_change(svc, ws):
    svc.sites.set_status(ws.site.id, "COMPLETED")
    with pytest.raises(ValidationError, match="can no longer be changed"):
        svc.materials.set_required(ws.site.id, [(ws.oak.id, "1")])


# --- supplier, bin and stock filters -----------------------------------------------------------------

def test_bin_filters_handle_products_in_several_bins(svc, ws):
    assert set(codes(svc, F(bin_id=ws.bin2.id))) == {"H3311"}
    assert codes(svc, F(bin_state="MULTIPLE")) == ["H3311"]
    assert "U708" in codes(svc, F(bin_state="NONE"))  # no default bin, no stock
    assert "H3311" not in codes(svc, F(bin_state="NONE"))
    # an archived bin is asked for its history
    svc.stock.transfer(ws.oak.id, ws.bin2.id, ws.bin1.id, "42")
    plan = svc.locations.remove_bin(ws.bin2.id, "rack removed")
    assert not plan.permanent  # it has history, so it is archived, not deleted
    assert codes(svc, F(bin_id=ws.bin2.id)) == ["H3311"]


def test_stock_filters_use_totals_over_every_bin(svc, ws):
    # 85 + 42 = 127 on hand; the reorder level (40) is judged on the total, not on one bin.
    svc.reservations.reserve({"product_id": ws.oak.id, "site_id": ws.site.id, "quantity": "100"})
    assert "H3311" in codes(svc, F(stock="LOW_STOCK"))  # 127 - 100 = 27 available <= 40
    assert "H3311" in codes(svc, F(stock="IN_STOCK"))
    assert "H3311" in codes(svc, F(stock="RESERVED"))
    assert "U708" in codes(svc, F(stock="OUT_OF_STOCK"))
    assert "U708" not in codes(svc, F(stock="AVAILABLE"))
    row = svc.materials.page(F(text="H3311"))[0]
    assert (row.on_hand, row.reserved, row.available, row.status_key) == (D(127), D(100), D(27), "LOW")
    assert [b.code for b in row.bins] == ["B04", "B05"]


def test_incoming_filter_matches_open_purchase_orders(svc, ws):
    from wms.services.purchasing import POLineInput
    po = svc.purchasing.save(None, {"supplier_id": ws.egger.id, "order_date": TODAY},
                             [POLineInput(ws.grey.id, "50")])
    svc.purchasing.mark_ordered(po.id)
    assert codes(svc, F(stock="INCOMING")) == ["U708"]
    assert svc.materials.page(F(stock="INCOMING"))[0].incoming == D(50)


def test_filters_combine_and_count_always_matches_the_rows(svc, ws):
    svc.materials.set_required(ws.site.id, [(ws.oak.id, "50"), (ws.grey.id, "25"), (ws.mdf.id, "3")])
    f = F(site_id=ws.site.id, supplier_id=ws.egger.id, bin_id=ws.bin1.id, stock="IN_STOCK", text="oak")
    assert codes(svc, f) == ["H3311"]
    for g in (f, F(site_id=ws.site.id), F(supplier_id=ws.egger.id), F(text="oak", stock="OUT_OF_STOCK"), F()):
        assert svc.materials.count(g.copy()) == len(codes(svc, g))


def test_sorting_and_paging(svc, ws):
    assert codes(svc, F(sort="qty", descending=True))[:2] == ["H3311", "MDF-18"]
    assert codes(svc, F(sort="name")) == sorted(codes(svc, F()), key=lambda c: {
        "H3311": "halifax", "U708": "light", "MDF-18": "mdf", "EDGE-W22": "abs", "CAB-600": "base"}[c])
    everything = codes(svc, F())
    pages = [r.product.internal_code for o in range(0, 5, 2) for r in svc.materials.page(F(), o, 2)]
    assert pages == everything  # no row lost or repeated between pages
    assert codes(svc, F(site_id=None, sort="net")) == everything  # project sort without project: ignored


def test_archived_status_filter(svc, ws):
    archive(svc, ws, ws.grey)
    assert "U708" not in codes(svc, F())
    assert codes(svc, F(status="ARCHIVED")) == ["U708"]
    assert "U708" in codes(svc, F(status="ALL"))


# --- export uses the same query -----------------------------------------------------------------------

def _sheet_codes(path, header_title="Code"):
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and r[0] == header_title)
    assert ws.freeze_panes == f"A{head + 2}" and ws.auto_filter.ref.startswith(f"A{head + 1}:")
    return [r[0] for r in rows[head + 1:]], rows[head]


def test_export_current_view_selected_and_all(svc, ws, tmp_path):
    f = F(supplier_id=ws.egger.id, sort="code", descending=True)
    n = export_materials(svc.materials, tmp_path / "view.xlsx", CURRENT, f)
    got, header = _sheet_codes(tmp_path / "view.xlsx")
    assert n == 2 and got == codes(svc, f) == ["U708", "H3311"]
    assert "Product Number" in header and "Available" in header
    sel = [ws.mdf.id, ws.oak.id]
    export_materials(svc.materials, tmp_path / "sel.xlsx", SELECTED, f, ids=sel)
    assert _sheet_codes(tmp_path / "sel.xlsx")[0] == ["MDF-18", "H3311"]
    archive(svc, ws, ws.grey)
    export_materials(svc.materials, tmp_path / "all.xlsx", ALL)
    assert set(_sheet_codes(tmp_path / "all.xlsx")[0]) == {"H3311", "U708", "MDF-18", "EDGE-W22", "CAB-600"}
    # quantities are numbers, not text
    wb = load_workbook(tmp_path / "view.xlsx").active
    values = [c.value for row in wb.iter_rows() for c in row]
    assert 127 in values


def test_export_project(svc, ws, tmp_path):
    svc.materials.set_required(ws.site.id, [(ws.oak.id, "50")])
    issue(svc, ws.site, ws.oak, 20, ws.bin1)
    export_project(svc.materials, ws.site.id, tmp_path / "p.xlsx")
    got, header = _sheet_codes(tmp_path / "p.xlsx", "Internal Code")
    assert got == ["H3311"]
    assert list(header[:10]) == ["Internal Code", "Material", "Product Number", "Supplier", "Unit", "Required",
                                 "Reserved", "Issued", "Returned", "Net Issued"]
    row = next(r for r in load_workbook(tmp_path / "p.xlsx").active.iter_rows(values_only=True) if r[0] == "H3311")
    assert row[5:10] == (50, 0, 20, 0, 20)
    assert "Project: P-1001" in str(load_workbook(tmp_path / "p.xlsx").active["A2"].value)


# --- saved views and bulk changes ---------------------------------------------------------------------

def test_saved_views_are_kept_apart_from_search_page(svc, ws):
    names = [v.name for v in svc.materials.views()]
    assert names[:3] == ["All Materials", "Low Stock", "No Bin"]
    svc.materials.save_view("Tribeca", F(site_id=ws.site.id, supplier_id=ws.egger.id))
    mine = [v for v in svc.materials.views() if v.name == "Tribeca"][0]
    assert (mine.filter.site_id, mine.filter.supplier_id) == (ws.site.id, ws.egger.id)
    assert "Tribeca" not in [s.name for s in svc.search.saved_searches()]
    with pytest.raises(ValidationError):
        svc.materials.save_view("Low Stock", F())
    svc.materials.delete_view(mine.id)
    assert "Tribeca" not in [v.name for v in svc.materials.views()]


def test_bulk_change_supplier_and_category(svc, ws):
    n = svc.materials.bulk_set([ws.mdf.id, ws.edge.id, ws.oak.id], "supplier_id", ws.egger.id)
    assert n == 2  # oak already had EGGER
    assert set(codes(svc, F(supplier_id=ws.egger.id))) == {"H3311", "U708", "MDF-18", "EDGE-W22"}
    assert svc.products.row(ws.mdf.id).on_hand == D(20)  # master data only
    with pytest.raises(ValidationError):
        svc.materials.bulk_set([ws.mdf.id], "name", 1)


def test_chips_describe_and_clear_each_filter(svc, ws):
    f = F(site_id=ws.site.id, supplier_id=ws.egger.id, bin_id=ws.bin1.id, stock="IN_STOCK", text="oak")
    chips = svc.materials.describe(f)
    assert [t for _, t in chips] == ["Project: Riverside Apartments", "Supplier: EGGER", "Bin: B04", "In Stock",
                                     "Search: oak"]
    svc.materials.clear(f, "supplier_id")
    assert f.supplier_id is None and f.site_id == ws.site.id


def test_no_artificial_row_limit(svc, ws):
    units = {u.code: u for u in svc.reference.units()}
    for i in range(620):
        svc.products.save(None, {"item_type": "RAW_MATERIAL", "internal_code": f"BULK-{i:04d}",
                                 "name": f"Bulk {i}", "unit_id": units["EA"].id})
    f = F(text="BULK")
    assert svc.materials.count(f) == 620
    got = [r for o in range(0, 620, 200) for r in svc.materials.page(f.copy(), o, 200)]
    assert len(got) == 620 and len({r.id for r in got}) == 620
