"""UI smoke tests: every screen opens, every record dialog opens, and the main
operator flows (Goods In, Goods Out) work through the real screens.

Message boxes are answered automatically (Yes / OK) and recorded, so an
unexpected error dialog fails the test.
"""

from __future__ import annotations

from decimal import Decimal

import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QDialog, QInputDialog, QMessageBox

from wms.ui.app import App


@pytest.fixture
def ui(qtbot, svc, monkeypatch):
    msgs = []

    def record(kind):
        def fn(parent, title, text, *a, **k):
            msgs.append((kind, title, text))
            return QMessageBox.StandardButton.Ok
        return fn

    monkeypatch.setattr(QMessageBox, "warning", record("warning"))
    monkeypatch.setattr(QMessageBox, "critical", record("critical"))
    monkeypatch.setattr(QMessageBox, "information", record("info"))

    def auto_yes(box):
        msgs.append(("confirm", box.windowTitle(), box.text()))
        for b in box.buttons():
            if box.buttonRole(b) == QMessageBox.ButtonRole.AcceptRole:
                b.click()
        return 0

    monkeypatch.setattr(QMessageBox, "exec", auto_yes)
    monkeypatch.setattr(QInputDialog, "getMultiLineText", lambda *a, **k: ("Test reason", True))
    opened = []

    def fake_exec(dlg):
        opened.append(type(dlg).__name__)
        dlg.show()
        dlg.close()
        return 0

    monkeypatch.setattr(QDialog, "exec", fake_exec)
    app = App(svc.ctx)
    from wms.ui.main_window import MainWindow
    win = MainWindow(app)
    qtbot.addWidget(win)
    win.show()
    return app, win, msgs, opened


def _errors(msgs):
    return [m for m in msgs if m[0] in ("warning", "critical")]


def test_every_page_opens(ui, sample):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 40)
    keys = list(win.factories)
    assert len(keys) == 20
    for key in keys:
        win.go(key)
        page = win.page(key)
        assert page is not None, key
        page.refresh()
    assert not _errors(msgs), msgs


def test_record_dialogs_open(ui, sample):
    app, win, msgs, opened = ui
    dn = sample.receive(sample.mdf, 40)
    from wms.services.goods_out import IssueLine
    from wms.services.production import MaterialInput
    from wms.services.stock_ops import ReturnLineInput
    gi = app.services.goods_out.post({"site_id": sample.site.id, "issue_date": dn.delivery_date},
                                     [IssueLine(sample.mdf.id, Decimal("5"))])
    rt = app.services.stock.post_return({"return_type": "FROM_SITE", "return_date": dn.delivery_date,
                                         "site_id": sample.site.id}, [ReturnLineInput(sample.mdf.id, "1",
                                                                                      sample.bin1.id)])
    po = app.services.production.post({"finished_product_id": sample.cabinet.id, "quantity": "1",
                                       "production_date": dn.delivery_date},
                                      [MaterialInput(sample.mdf.id, "2", None)])
    for kind, oid in (("PRODUCT", sample.mdf.id), ("PRODUCT", sample.cabinet.id), ("DELIVERY_NOTE", dn.id),
                      ("GOODS_ISSUE", gi.id), ("SITE", sample.site.id), ("RETURN", rt.id), ("BIN", sample.bin1.id),
                      ("SUPPLIER", sample.supplier.id), ("PRODUCTION", po.id)):
        app.open_entity(kind, oid, win)
    from wms.ui.dialogs.stock_dialogs import AdjustmentDialog, ProductionEntryDialog, ReturnEntryDialog, TransferDialog
    from wms.ui.pages.movements import MovementDetailDialog
    TransferDialog(app, sample.mdf.id, win, sample.bin1.id)
    AdjustmentDialog(app, sample.mdf.id, sample.bin1.id, win)
    ReturnEntryDialog(app, "TO_SUPPLIER", win)
    ProductionEntryDialog(app, win)
    MovementDetailDialog(app, app.services.stock.movements()[0], win)
    assert len(opened) == 9
    assert not _errors(msgs), msgs


def test_goods_in_screen_posts_delivery(ui, sample):
    app, win, msgs, _ = ui
    win.go("goods_in")
    page = win.page("goods_in")
    page.supplier.set_value(sample.supplier.id)
    page.dn_number.setText("DN 7781")
    page.lines.clear()
    page.lines.add_row(sample.mdf.id, Decimal("12"))
    page.lines.add_row(sample.edge.id, Decimal("150.5"), sample.bin1.id)
    page.post()
    assert not _errors(msgs), msgs
    dns = app.services.goods_in.search("DN 7781")
    assert len(dns) == 1 and dns[0][1] == 2
    assert page.dn_number.text() == ""  # form cleared after posting
    # Entering the same delivery note again is blocked
    page.supplier.set_value(sample.supplier.id)
    page.dn_number.setText("dn7781")
    page.lines.add_row(sample.mdf.id, Decimal("1"))
    page.post()
    assert any(t == "Already entered" for _, t, _ in msgs)
    assert len(app.services.goods_in.search("7781")) == 1


def test_goods_out_screen_blocks_then_overrides_negative(ui, sample):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 3)
    win.go("goods_out")
    page = win.page("goods_out")
    page.site.set_value(sample.site.id)
    page.lines.clear()
    page.lines.add_row(sample.mdf.id, Decimal("5"))
    page.post()  # admin may override: confirm + reason are auto-answered
    assert any(t == "Not enough stock" for _, t, _ in msgs)
    issues = app.services.goods_out.search()
    assert len(issues) == 1
    assert app.services.products.get(sample.mdf.id) is not None
    rows = {r.product.id: r for r in app.services.products.list()}
    assert rows[sample.mdf.id].on_hand == Decimal("-2")


def test_goods_out_double_post_and_repeat_are_caught(ui, sample):
    """Ctrl+S pressed twice / the same issue typed in twice must not silently create two issues."""
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 20)
    win.go("goods_out")
    page = win.page("goods_out")
    page.site.set_value(sample.site.id)
    page.lines.clear()
    page.lines.add_row(sample.mdf.id, Decimal("2"))
    token = page._token
    page.post()
    assert len(app.services.goods_out.search()) == 1 and page._token != token  # new form, new token
    # The server refuses the old form's token even if the screen somehow sent it again
    from wms.core.exceptions import AlreadyPostedError
    from wms.services.goods_out import IssueLine
    with pytest.raises(AlreadyPostedError):
        app.services.goods_out.post({"site_id": sample.site.id, "issue_date": page.date.value(),
                                     "client_token": token}, [IssueLine(sample.mdf.id, "2")])
    # Re-entering the same materials for the same site asks first (auto-answered "Post again" here)
    page.site.set_value(sample.site.id)
    page.lines.clear()
    page.lines.add_row(sample.mdf.id, Decimal("2"))
    page.post()
    assert any(t == "Possible duplicate" for _, t, _ in msgs)
    assert len(app.services.goods_out.search()) == 2
    # Re-entrancy: a post() arriving while one is running is ignored
    page._posting = True
    page.post()
    assert len(app.services.goods_out.search()) == 2
    assert not _errors(msgs), msgs


def test_scanner_flow_in_lines_editor(ui, sample):
    """A USB scanner types the barcode and presses Enter: product chosen, cursor to quantity, Enter -> next line."""
    app, win, msgs, _ = ui
    app.services.products.save(sample.mdf.id, {"item_type": "RAW_MATERIAL", "internal_code": "MDF-18",
                                               "name": "MDF 18mm 2440x1220", "unit_id": sample.mdf.unit_id,
                                               "barcode": "5012345678900", "default_bin_id": sample.bin1.id})
    win.go("goods_in")
    page = win.page("goods_in")
    page.refresh()
    ed = page.lines
    ed.clear()
    ed.add_row()
    prod = ed.table.cellWidget(0, 0)
    prod.setEditText("5012345678900")
    prod.lineEdit().returnPressed.emit()
    assert prod.value() == sample.mdf.id
    # Next: the bin (blank = the product's default bin); scanning a bin label fills it
    bin_w = ed._bin_widget(0)
    assert bin_w.value() is None
    bin_w.setEditText("b05")
    bin_w.lineEdit().returnPressed.emit()
    assert bin_w.value() == sample.bin2.id
    bin_w.setEditText("NO-SUCH-BIN")
    bin_w.lineEdit().returnPressed.emit()
    assert bin_w.value() is None
    bin_w.setEditText("B04")
    bin_w.lineEdit().returnPressed.emit()
    assert bin_w.value() == sample.bin1.id
    qty = ed.table.cellWidget(0, 1)
    qty.setText("4")
    qty.returnPressed.emit()
    assert ed.table.rowCount() == 2  # ready for the next scan
    unknown = ed.table.cellWidget(1, 0)
    unknown.setEditText("NO-SUCH-CODE")
    unknown.lineEdit().returnPressed.emit()
    assert unknown.value() is None and "No product" in ed.total_label.text()
    assert [(v.product_id, v.quantity) for v in ed.values()] == [(sample.mdf.id, Decimal("4"))]


def test_enter_in_a_dialog_field_never_presses_a_button(ui, sample):
    from PySide6.QtWidgets import QPushButton

    from wms.ui.dialogs.stock_dialogs import ReturnEntryDialog, TransferDialog
    app, *_ = ui
    for dlg in (TransferDialog(app), ReturnEntryDialog(app)):
        assert all(not b.autoDefault() and not b.isDefault() for b in dlg.findChildren(QPushButton))


def test_stocktake_screen_count_and_approve(ui, sample):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 10)
    count = app.services.stocktake.create(None)
    win.go("stocktake")
    page = win.page("stocktake")
    page.refresh()
    page._show(app.services.stocktake.get(count.id))
    page.lines.cellWidget(0, 5).setText("8")
    page.lines.cellWidget(0, 7).setText("2 sheets water damaged")
    page.submit()
    assert app.services.stocktake.get(count.id).status == "SUBMITTED"
    page.approve()
    assert app.services.stocktake.get(count.id).status == "APPROVED"
    assert {r.product.id: r for r in app.services.products.list()}[sample.mdf.id].on_hand == Decimal("8")
    assert not _errors(msgs), msgs


def test_dashboard_shows_actionable_alerts(ui, sample):
    from datetime import date

    from wms.services.goods_out import IssueLine
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 1)
    app.services.goods_out.post({"site_id": sample.site.id, "issue_date": date.today()},
                                [IssueLine(sample.mdf.id, "3")], override_reason="sheets on site")
    d = app.services.dashboard.load()
    texts = " | ".join(a.text for a in d.alerts)
    assert "negative" in texts and "no scan" in texts
    low = d.action("low_stock")
    assert low.count >= 1 and "out of stock" in low.hint and low.level == "danger"
    assert d.today["Goods Out"] == 1 and d.today["Goods In"] == 1
    win.go("dashboard")
    dash = win.page("dashboard")
    dash.refresh()
    assert dash.tiles["low_stock"].objectName() == "TileDanger"
    assert dash.tiles["no_bin"].text().endswith("Products without bin")
    # Every counter opens the list of exactly those records
    dash.tiles["low_stock"].click()
    assert win.page("low_stock").isVisible()
    win.go("dashboard")
    dash.open_action("no_photo")
    search = win.page("search")
    assert search.criteria.photo == "WITHOUT"
    win.go("dashboard")
    dash.open_action("overdue_po")
    assert win.page("purchase_orders").status.currentData() == "OVERDUE"
    assert not _errors(msgs), msgs


def test_operative_sees_limited_menu(qtbot, svc, sample):
    roles = {r.code: r for r in svc.users.list_roles()}
    svc.users.create_user({"username": "op1", "full_name": "Op One", "role_id": roles["OPERATIVE"].id,
                           "must_change_password": False}, "secret123")
    svc.auth.login("op1", "secret123")
    from wms.ui.main_window import MainWindow
    win = MainWindow(App(svc.ctx))
    qtbot.addWidget(win)
    assert "goods_in" in win.factories and "goods_out" in win.factories
    assert "users" not in win.factories and "audit" not in win.factories




def test_theme_ignores_windows_dark_mode(qapp):
    """Windows dark mode gives Qt white text; the app must still show dark text on light backgrounds."""
    from PySide6.QtGui import QColor, QPalette

    from wms.ui.theme import apply_theme
    old_pal, old_qss = qapp.palette(), qapp.styleSheet()
    dark = QPalette()
    dark.setColor(QPalette.ColorRole.WindowText, QColor("#FFFFFF"))
    dark.setColor(QPalette.ColorRole.Text, QColor("#FFFFFF"))
    qapp.setPalette(dark)
    try:
        apply_theme(qapp)
        pal = qapp.palette()
        for role in (QPalette.ColorRole.WindowText, QPalette.ColorRole.Text, QPalette.ColorRole.ButtonText):
            assert pal.color(role).lightness() < 100, role
        assert pal.color(QPalette.ColorRole.Base).lightness() > 240
    finally:
        qapp.setPalette(old_pal)
        qapp.setStyleSheet(old_qss)


# --- Search screen ---------------------------------------------------------------------------

def _search_page(win):
    win.go("search")
    return win.page("search")


def test_search_screen_type_results_filter_open(ui, sample):
    app, win, msgs, opened = ui
    sample.receive(sample.mdf, 40)
    page = _search_page(win)
    page.search_for("mdf")
    view = page.products_view
    assert view.count.text() == "1 product" and view.table.items()[0].product.id == sample.mdf.id
    assert page.tabs.tabText(0) == "Products (1)"
    row = view.table.items()[0]
    assert (row.location_text, row.bin_text) == ("MAIN / A / R03 / S02", "B04")
    page.search_for("B04")
    assert page.side.isVisible() and page.matches.item(0).text() == "BINS"  # group heading
    assert page.matches.item(1).text().startswith("B04\n")
    page._open_match(page.matches.item(0))  # a heading opens nothing
    page._open_match(page.matches.item(1))
    assert opened[-1] == "BinContentsDialog"
    # Nothing found: an explicit message instead of an empty table
    page.search_for("no-such-thing-zz")
    assert view.stack.currentWidget() is view.empty and view.count.text() == "No matching records found"
    assert "No matching records found" in view.empty.text()
    assert "zz" in app.services.search.recent()[0]
    assert not _errors(msgs), msgs


def test_search_screen_filters_chips_and_clear_all(ui, sample):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 40)
    page = _search_page(win)
    page.filters_btn.click()
    assert page.panel.isVisible()
    page.panel.supplier.set_value(sample.supplier.id)
    page.panel.stock_status.setCurrentIndex(page.panel.stock_status.findData("IN_STOCK"))
    assert page.criteria.supplier_id == sample.supplier.id and page.criteria.stock_status == "IN_STOCK"
    chips = [page.chips.itemAt(i).widget().text() for i in range(page.chips.count())
             if page.chips.itemAt(i).widget() is not None]
    assert any(t.startswith("Supplier: Egger UK Ltd") for t in chips) and "Clear All Filters" in chips
    assert page.products_view.count.text() == "1 product"
    page.remove_filter("supplier_id")  # the x on one chip
    assert page.criteria.supplier_id is None and page.panel.supplier.value() is None
    assert page.criteria.stock_status == "IN_STOCK"
    page.clear_filters()
    assert not page.criteria.has_filters and not page.chips_w.isVisible()
    assert page.products_view.count.text() == "3 products"
    assert not _errors(msgs), msgs


def test_search_screen_quick_filters_sorting_and_transactions(ui, sample):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 4)  # below its minimum of 10
    page = _search_page(win)
    low = next(b for b in page.quick.buttons() if b.property("quick") == "LOW_STOCK")
    low.click()
    assert [r.product.internal_code for r in page.products_view.table.items()] == ["MDF-18"]
    next(b for b in page.quick.buttons() if b.property("quick") == "ALL").click()
    table = page.products_view.table
    qty_col = next(i for i, c in enumerate(table.model_.columns) if c.sort == "qty")
    table.horizontalHeader().sectionClicked.emit(qty_col)
    assert page.sorts[page.PRODUCTS] == ("qty", True)
    title = table.model_.headerData(qty_col, Qt.Orientation.Horizontal)
    assert title.endswith("▼"), title  # biggest first, and the header says so
    assert [r.product.internal_code for r in table.items()][0] == "MDF-18"
    page.tabs.setCurrentWidget(page.movements_view)
    assert page.movements_view.count.text() == "1 transaction"
    assert not _errors(msgs), msgs


def test_scan_into_search_then_goods_out_in_one_click(ui, sample):
    app, win, msgs, _ = ui
    app.services.products.save(sample.mdf.id, {"item_type": "RAW_MATERIAL", "internal_code": "MDF-18",
                                               "name": "MDF 18mm 2440x1220", "unit_id": sample.mdf.unit_id,
                                               "barcode": "5012345678900", "default_bin_id": sample.bin1.id,
                                               "supplier_id": sample.supplier.id})
    sample.receive(sample.mdf, 10)
    win.search.setText("5012345678900")  # scanner into the top bar: code + Enter
    win.global_search()
    page = win.page("search")
    assert win.stack.currentWidget() is page and page.box.text() == "5012345678900"
    assert page.exact_w.isVisible() and "MDF-18" in page.exact_label.text()
    page._act_exact("goods_out")
    out = win.page("goods_out")
    assert win.stack.currentWidget() is out
    assert [v.product_id for v in out.lines.values(validate=False) if v.product_id] == [sample.mdf.id]
    win.go("search")
    page._act_exact("goods_in")
    gin = win.page("goods_in")
    assert [v.product_id for v in gin.lines.values(validate=False) if v.product_id] == [sample.mdf.id]
    assert gin.supplier.value() == sample.supplier.id
    assert not _errors(msgs), msgs


def test_saved_search_reopens_with_one_click(ui, sample, monkeypatch):
    app, win, msgs, _ = ui
    page = _search_page(win)
    page.filters_btn.click()
    page.panel.supplier.set_value(sample.supplier.id)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("Egger stock", True))
    page.save_search()
    saved = app.services.search.saved_searches()
    assert [s.name for s in saved] == ["Egger stock"]
    page.clear_filters()
    page.open_saved(saved[0])
    assert page.criteria.supplier_id == sample.supplier.id and page.panel.supplier.value() == sample.supplier.id
    assert not _errors(msgs), msgs


def test_search_result_actions_open_the_right_screens(ui, sample):
    app, win, msgs, opened = ui
    sample.receive(sample.mdf, 10)
    page = _search_page(win)
    page.search_for("MDF-18")
    for action in ("view", "history", "delivery_note", "transfer"):
        page.product_action(action)
    assert opened[-4:] == ["ProductDialog", "ProductDialog", "DeliveryNoteDialog", "TransferDialog"]
    page.product_action("photo")  # no photo yet: a message, not an error
    assert msgs[-1][0] == "info"
    assert not _errors(msgs), msgs


def test_archived_records_open_with_their_history_and_restore(ui, sample, monkeypatch):
    """Delete B05 through the screen (it has history, so it is archived), open its history page,
    then restore it - all without error dialogs."""
    app, win, msgs, opened = ui
    sample.receive(sample.mdf, 3, sample.bin2)
    app.services.stock.transfer(sample.mdf.id, sample.bin2.id, sample.bin1.id, "3")
    from wms.ui.dialogs import removal
    from wms.ui.dialogs.documents import BinContentsDialog
    monkeypatch.setattr(removal.RemovalDialog, "exec", lambda self: QDialog.DialogCode.Accepted)
    dlg = BinContentsDialog(app, sample.bin2.id, win)
    dlg.delete_bin()
    assert not app.services.locations.reload(type(sample.bin2), sample.bin2.id).is_active
    assert "Archived" in dlg.title.text() and "Historical movements: 2" in dlg.summary.text()
    assert dlg.history.model().rowCount() == 2 and dlg.archived.isVisibleTo(dlg)
    assert any(m[1] == "Removed from active use" for m in msgs)
    app.services.sites.set_status(sample.site.id, "COMPLETED")
    app.open_entity("SITE", sample.site.id, win)
    app.services.locations.restore_bin(sample.bin2.id)
    dlg.refresh()
    assert "Archived" not in dlg.title.text()
    assert not _errors(msgs), msgs


def test_reverse_a_transfer_from_the_movement_screen(ui, sample, monkeypatch):
    """Movement detail -> Reverse... -> pick a reason -> mirror movement posted and shown as linked."""
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 10, sample.bin1)
    tr = app.services.stock.transfer(sample.mdf.id, sample.bin1.id, sample.bin2.id, "4")
    from wms.ui.pages import movements
    asked = []

    def pick(parent, title, prompt, choices):
        asked.append([c for c, _ in choices])
        return "DATA_ENTRY_ERROR", "Scanned the wrong bin"

    monkeypatch.setattr(movements, "ask_reason", pick)
    dlg = movements.MovementDetailDialog(app, app.services.stock.get_movement(tr.movement_id))
    dlg.reverse()
    assert "DATA_ENTRY_ERROR" in asked[0] and "SCRAP" not in asked[0]  # only reversal reasons offered
    rec = app.services.ledger.reversal_of(tr.movement_id)
    assert rec.reason == "Scanned the wrong bin"
    again = movements.MovementDetailDialog(app, app.services.stock.get_movement(rec.reversal_movement_id))
    assert again.mv.type_label == "Reversal of Bin Transfer"
    win.go("movements")
    page = win.page("movements")
    idx = page.mtype.findData("REVERSAL")
    page.mtype.setCurrentIndex(idx)
    page.refresh()
    assert [m.id for m in page.table.items()] == [rec.reversal_movement_id]
    assert not _errors(msgs), msgs


def test_reason_codes_tab_lists_defaults(ui):
    app, win, msgs, _ = ui
    win.go("settings")
    page = win.page("settings")
    page.refresh()
    codes = {r.code for r in page.reason_table.items()}
    assert {"DAMAGED", "MISSING", "STOCK_COUNT", "SCRAP", "OTHER"} <= codes
    assert not _errors(msgs), msgs


def test_bin_screen_shows_location_levels_and_prints_a_label(ui, sample, monkeypatch):
    app, win, msgs, _ = ui
    sample.receive(sample.mdf, 6, sample.bin1)
    from wms.ui.dialogs import documents
    opened = []
    monkeypatch.setattr(documents, "open_file", lambda path: opened.append(str(path)))
    dlg = documents.BinContentsDialog(app, sample.bin1.id)
    text = dlg.where.text()
    assert "Zone: Zone A" in text and "Rack: R03" in text and "Status: Active" in text
    assert f"Barcode: {sample.bin1.full_code}" in text
    documents.print_bin_labels(dlg, app, [sample.bin1.id])
    assert opened and opened[0].endswith(".pdf")
    dlg.stocktake()
    assert app.services.stocktake.list()[0].lines[0].bin_id == sample.bin1.id
    dlg.close()
    dlg.deleteLater()
    assert not _errors(msgs), msgs


def test_purchase_order_to_partial_delivery_and_problem(ui, sample, qtbot):
    app, win, msgs, opened = ui
    from wms.ui.dialogs.purchasing import DiscrepancyDialog, NewDiscrepancyDialog, PurchaseOrderDialog
    dlg = PurchaseOrderDialog(app, None)
    qtbot.addWidget(dlg)
    dlg.supplier.set_value(sample.supplier.id)
    dlg.lines.clear()
    dlg.lines.add_row(sample.mdf.id, Decimal("10"))
    assert dlg.save(), msgs
    po_id = dlg.po_id
    dlg.mark_ordered()
    assert app.services.purchasing.get(po_id).po.status == "ORDERED"
    win.go("purchase_orders")
    page = win.page("purchase_orders")
    page.refresh()
    assert [r.po.id for r in page.table.items()] == [po_id]

    page.receive(po_id)  # Goods In opens with the outstanding lines
    gin = win.page("goods_in")
    assert gin.order.value() == po_id and gin.supplier.value() == sample.supplier.id
    assert [(v.product_id, v.quantity) for v in gin.lines.values()] == [(sample.mdf.id, Decimal("10"))]
    gin.lines.table.cellWidget(0, 1).set_value(Decimal("6"))  # only 6 arrived
    gin.dn_number.setText("DN-500")
    gin.post()
    assert not _errors(msgs), msgs
    assert "partially received" in msgs[-1][2]
    v = app.services.purchasing.get(po_id)
    assert v.po.status == "PARTIALLY_RECEIVED" and v.lines[0].outstanding == Decimal("4")
    dn_id = v.deliveries[0].id

    new = NewDiscrepancyDialog(app, None, delivery_note_id=dn_id)
    qtbot.addWidget(new)
    new.kind.setCurrentIndex(new.kind.findData("DAMAGED"))
    new.product.set_value(sample.mdf.id)
    new.qty.set_value(Decimal("1"))
    new.save()
    d = new.created
    assert d.purchase_order_id == po_id and d.reason_code == "DAMAGED"
    resolve = DiscrepancyDialog(app, d.id)
    qtbot.addWidget(resolve)
    resolve.resolve()
    assert app.services.purchasing.get_discrepancy(d.id).status == "RESOLVED"

    for kind, oid in (("PURCHASE_ORDER", po_id), ("DISCREPANCY", d.id), ("DELIVERY_NOTE", dn_id)):
        app.open_entity(kind, oid)
    closing = PurchaseOrderDialog(app, po_id)
    qtbot.addWidget(closing)
    closing.close_short()
    assert app.services.purchasing.get(po_id).po.closed_short
    win.go("discrepancies")
    win.page("discrepancies").refresh()
    win.go("low_stock")
    win.page("low_stock").refresh()
    assert not _errors(msgs), msgs


def test_reserve_for_project_then_reorder_required(ui, sample, qtbot):
    """Stage 4 on screen: reserve on the product card, see the stock states, get warned when another
    project's reservation is about to be issued, and turn the Reorder Required list into draft orders."""
    app, win, msgs, opened = ui
    from wms.ui.dialogs.product_dialog import ProductDialog
    from wms.ui.dialogs.reservations import ReservationDialog
    sample.receive(sample.mdf, 40)
    other = app.services.sites.save(None, {"project_number": "P-2002", "name": "Hill House", "status": "ACTIVE"})

    dlg = ReservationDialog(app, None, product_id=sample.mdf.id, site_id=sample.site.id)
    qtbot.addWidget(dlg)
    assert "available 40" in dlg.states.text()
    dlg.qty.set_value(Decimal("35"))
    dlg.save()
    res = dlg.saved.reservation
    assert res.number.startswith("RES-") and not _errors(msgs)

    too_much = ReservationDialog(app, None, product_id=sample.mdf.id, site_id=other.id)
    qtbot.addWidget(too_much)
    too_much.qty.set_value(Decimal("10"))
    too_much.save()
    assert any("Only 5" in text for _, _, text in _errors(msgs))  # over-reservation refused
    msgs.clear()

    card = ProductDialog(app, sample.mdf.id)
    qtbot.addWidget(card)
    states = {k: w.text() for k, w in card.ov_states.items()}
    assert states == {"on_hand": "40 SHEET", "reserved": "35 SHEET", "available": "5 SHEET",
                      "incoming": "0 SHEET", "minimum": "10 SHEET"}
    assert "Reorder" in card.ov_reorder.text() or "Low" in card.ov_reorder.text()
    assert [v.reservation.id for v in card.reservations.table.items()] == [res.id]
    card.reorder_level.set_value(Decimal("12"))
    card.save()
    assert app.services.products.get(sample.mdf.id).reorder_level == Decimal("12")

    # Goods Out to the other project would eat P-1001's reservation: warned first (auto-confirmed here).
    win.go("goods_out")
    page = win.page("goods_out")
    page.site.set_value(other.id)
    page.lines.clear()
    page.lines.add_row(sample.mdf.id, Decimal("8"))
    page.post()
    warning = [text for _, title, text in msgs if title == "Reserved for another project"]
    assert warning and "P-1001" in warning[0] and "uses 3" in warning[0]

    win.go("reservations")
    rp = win.page("reservations")
    rp.refresh()
    assert [v.reservation.id for v in rp.table.items()] == [res.id]
    app.open_entity("RESERVATION", res.id)
    app.open_entity("SITE", sample.site.id, tab="reservations")

    win.go("low_stock")
    low = win.page("low_stock")
    low.refresh()
    assert [r.product.id for r in low.table.items()] == [sample.mdf.id]
    low.draft_orders()
    drafts = app.services.purchasing.list()
    assert len(drafts) == 1 and drafts[0].po.status == "DRAFT"
    assert any(title == "Draft orders" and drafts[0].po.po_number in text for _, title, text in msgs)
    assert not _errors(msgs), msgs


def test_scan_transfer_and_goods_out_project(ui, sample, qtbot):
    """Transfer: scan product, FROM bin, TO bin, quantity, Ctrl+S. Goods Out: scan the project number."""
    app, win, msgs, _ = ui
    from wms.ui.dialogs.stock_dialogs import TransferDialog
    sample.receive(sample.mdf, 12)
    dlg = TransferDialog(app)
    qtbot.addWidget(dlg)
    for field, scan in ((dlg.product, "mdf-18"), (dlg.src, "B04"), (dlg.dst, "B05")):
        field.setEditText(scan)
        field.lineEdit().returnPressed.emit()
    assert (dlg.product.value(), dlg.src.value(), dlg.dst.value()) == (sample.mdf.id, sample.bin1.id, sample.bin2.id)
    dlg.qty.setText("5")
    dlg.save()
    levels = {lvl.bin_id: lvl.quantity for lvl in app.services.products.stock_by_bin(sample.mdf.id)}
    assert levels == {sample.bin1.id: Decimal("7"), sample.bin2.id: Decimal("5")}

    win.go("goods_out")
    page = win.page("goods_out")
    page.refresh()
    page.site.setEditText("p-1001")
    page.site.lineEdit().returnPressed.emit()
    assert page.site.value() == sample.site.id
    assert not _errors(msgs), msgs


def test_ready_made_views_menu(ui, sample):
    app, win, msgs, _ = ui
    from wms.services.search import BUILTIN_VIEWS
    page = _search_page(win)
    views = dict(BUILTIN_VIEWS)
    page.open_builtin(views["Products Without Bin"])
    assert page.criteria.default_bin == "WITHOUT" and page.chips_w.isVisibleTo(page)
    page.open_builtin(views["Recently Received"])
    assert page.criteria.quick == "RECENTLY_RECEIVED" and page.criteria.default_bin is None
    page.open_builtin(views["Outstanding Purchase Orders"])
    assert win.nav.currentItem().data(Qt.ItemDataRole.UserRole) == "purchase_orders"
    assert not _errors(msgs), msgs


def test_excel_import_wizard_and_exports(ui, sample, qtbot, tmp_path):
    """Stage 6 on screen: workbook -> worksheet -> columns -> preview (action per row) -> import with a
    backup first; then export all / filtered / selected products, search results and a product's history."""
    app, win, msgs, _ = ui
    from openpyxl import Workbook, load_workbook

    from wms.ui.dialogs.excel_import import ExcelImportDialog
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.append(["Code", "Name / Description", "Product Number", "Supplier", "Unit"])
    ws.append(["H3311", "Egger Halifax oak", "H3311 ST10", "Egger UK Ltd", "SHEET"])
    ws.append(["MDF-18", "MDF renamed by Excel", "", "", "SHEET"])
    ws.append(["U708", "Light grey", "", "New Boards Ltd", "BOXES"])
    wb.create_sheet("Notes")
    path = tmp_path / "materials.xlsx"
    wb.save(path)

    dlg = ExcelImportDialog(app, win)
    qtbot.addWidget(dlg)
    dlg.next()  # no file chosen
    assert dlg.stack.currentIndex() == 0 and _errors(msgs)
    msgs.clear()
    dlg.path.setText(str(path))
    dlg.next()
    assert dlg.stack.currentIndex() == 1 and dlg.sheet.currentData() == "Materials"
    assert dlg.sample.rowCount() == 3 and "No photos found" in dlg.sheet_summary.text()
    dlg.next()
    assert dlg.map_combos["internal_code"].currentData() == 0 and dlg.map_combos["name"].currentData() == 1
    assert dlg.map_combos["photo"].currentData() is None
    dlg.next()
    assert dlg.stack.currentIndex() == 3 and dlg.preview.rowCount() == 3
    assert "1 new, 0 to update, 1 skipped, 1 with errors" in dlg.preview_summary.text()
    # Choose "Update existing" for MDF-18 in the Action column
    action = dlg.preview.cellWidget(1, 1)
    action.setCurrentIndex(action.findData("UPDATE"))
    assert "1 new, 1 to update, 0 skipped" in dlg.preview_summary.text()
    dlg.only_problems.setChecked(True)
    assert dlg.preview.rowCount() == 1  # U708: unknown unit (and a new-supplier note)
    dlg.next()
    assert dlg.stack.currentIndex() == 4, msgs
    assert "Pre-import backup successfully created" in dlg.done_text.text()
    assert any(title == "Import from Excel" and "1 row(s) with errors will NOT be imported" in text
               for _, title, text in msgs)
    assert app.services.products.get(sample.mdf.id).name == "MDF renamed by Excel"
    assert app.services.search.exact_product("H3311").supplier.name == "Egger UK Ltd"
    assert app.services.search.exact_product("U708") is None

    win.go("products")
    page = win.page("products")
    page.refresh()
    assert page.export_excel("selected", str(tmp_path / "none.xlsx")) is None
    assert any("Select one or more" in text for _, _, text in msgs)
    page.search.setText("MDF")
    page.refresh()
    page.export_excel("shown", str(tmp_path / "shown.xlsx"))
    rows = [r[0].value for r in load_workbook(tmp_path / "shown.xlsx").active.iter_rows(min_row=6)]
    assert rows[:1] == ["MDF-18"] and "H3311" not in rows
    page.search.setText("")
    page.refresh()
    page.table.selectRow(0)
    page.table.selectionModel().select(page.table.proxy.index(2, 0),
                                       page.table.selectionModel().SelectionFlag.Select
                                       | page.table.selectionModel().SelectionFlag.Rows)
    assert len(page.table.selected_items()) == 2
    page.export_excel("selected", str(tmp_path / "sel.xlsx"))
    page.export_excel("all", str(tmp_path / "all.xlsx"))
    ws = load_workbook(tmp_path / "all.xlsx").active
    assert ws["A3"].value == "All products - 4 product(s)"
    assert len([r for r in ws.iter_rows(min_row=6) if r[0].value and r[0].value != "Total"]) == 4

    search = _search_page(win)
    search.search_for("egger")
    out = search.export_products(str(tmp_path / "search.xlsx"))
    assert out is not None
    codes = {r[0].value for r in load_workbook(out).active.iter_rows(min_row=6)}
    assert {"H3311", "MDF-18"} <= codes

    from wms.ui.dialogs.product_dialog import ProductDialog
    sample.receive(sample.mdf, 5)
    card = ProductDialog(app, sample.mdf.id)
    qtbot.addWidget(card)
    out = card.export_history(str(tmp_path / "history.xlsx"))
    ws = load_workbook(out).active
    assert ws.title == "History MDF-18" and ws["A5"].value == "Date / Time" and ws["C6"].value == "MDF-18"
    assert not _errors(msgs), msgs


def test_backup_health_schedule_and_test_restore_on_screen(ui, sample, cfg, tmp_path):
    """Stage 7 on screen: the Dashboard says when the last good backup was, the schedule makes one,
    a failure is shown (never silent), and Settings can test-restore a backup."""
    app, win, msgs, _ = ui
    win.go("dashboard")
    dash = win.page("dashboard")
    dash.refresh()
    assert dash.backup_line.text() == "Last successful backup: Never"
    assert "No backup has been made yet" in dash.t_alerts.items()[0].text
    path = win.scheduled_backup()
    assert path is not None and path.name.endswith("_auto.zip")
    dash.refresh()
    assert dash.backup_line.text().startswith("Last successful backup: Today")
    assert win.scheduled_backup() is None  # not due again yet

    cfg.backup_dir = str(tmp_path / "file-in-the-way" / "x")
    (tmp_path / "file-in-the-way").write_text("x")
    import json
    st = json.loads(app.services.backup.status_path.read_text())
    st["last_success_at"] = "2020-01-01T00:00:00"
    app.services.backup.status_path.write_text(json.dumps(st))
    assert win.scheduled_backup() is None
    assert any(title == "Automatic backup failed" for kind, title, _ in msgs if kind == "warning")
    msgs.clear()
    dash.refresh()
    assert "FAILED" in dash.t_alerts.items()[0].text

    cfg.backup_dir = str(tmp_path / "backups")
    win.go("settings")
    page = win.page("settings")
    page.refresh()
    assert "FAILED" in page.bk_health.text()
    good = app.services.backup.create(note="for the test restore")
    page.refresh()
    assert page.bk_health.text().startswith("Last successful backup: Today")
    page.bk_table.selectRow(0)
    assert page.bk_table.current().path == good
    result = page._test_restore()
    assert result is not None and result.ok
    assert any(title == "Test restore" and "OK - restorable" in text for _, title, text in msgs)
    assert not _errors(msgs), msgs
